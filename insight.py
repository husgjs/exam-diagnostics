# -*- coding: utf-8 -*-
"""教师工具 · 洞察三件套：难点预判 / 连续下滑预警 / 教学复盘

    python insight.py hardkp --kp 绝对值            难点预判（读题库 + 逐题得分）
    python insight.py drop --files 第一次.xlsx 第二次.xlsx [第三次...]   连续下滑预警
    python insight.py reflect --files 成绩单.xlsx [--scores 逐题.csv] [--class 九年级2班]

三条规矩（沿用错题画像 ask_bank 立下的那两条，再加一条）：
  1. 样本不足就明说不足，不编结论。题数不够、只有一次数据、缺逐题得分，都直接讲清楚。
  2. 所有判定阈值写成顶部常量，并在注释里写清依据；阈值不许散落在代码里。
  3. 教学复盘只谈知识点与讲法，**禁止用成绩数据评价老师个人**。输出前过一遍禁语自检
     （见 FORBID_WORDS / guard_text），命中就报错而不是悄悄写进报告。

数据质检：任何统计之前先质检（0 分/缺考先排除并提示），否则均分会被人为拉低。
模型调用：本工具 **0 次**模型调用 —— 所有数字由代码算出，所有文案都是本地模板。

输入格式（与 grade_reports.py / teaching.py 已有的口径对齐，不另立一套）
  成绩单 xlsx：第一张同时含「姓名」和「总分」（精确匹配，不用「小总分」）表头的 sheet；
               有「班级」列就用，有各科分（语文分数 等）就算科目录。
  逐题得分：  CSV/TXT（推荐，阅卷系统直接导出）或 xlsx，表头含
               学生,题号,知识点,满分,得分        # 知识点/满分可省，省了就用题库补
             也认「学生,题号,得分」三列，或宽表（姓名 + 每列一个题号，值必须是数字分）。
"""
import argparse
import csv
import json
import os
import sys
import datetime

try:
    if (sys.stdout.encoding or "").lower().replace("-", "") not in ("utf8", "utf_8"):
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

import openpyxl
import ingest
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT_DIR = os.path.join(HERE, "输出")
BANK = os.path.join(HERE, "题库.json")
DEFAULT_GRADES = os.path.join(ROOT, "九年级月考成绩单.xlsx")

# ══════════════════════════════════════════════════════════════════════════
# 判定阈值 —— 全部集中在这里。改口径只改这几行，代码里不许再出现魔法数字。
# ══════════════════════════════════════════════════════════════════════════

# 难点预判：知识点下至少要有几道题才敢下判断。
# 依据：1 道题说明不了「这个点的坎在哪」，2 道刚好可能一道送分一道超纲，
# 3 道才有「同一个点上的不同问法」可比较 —— 与五维雷达「样本 3 条起才评分」同一口径。
MIN_ITEMS = 3

# 单题得分率分档（得分率 P = 该题平均得分 / 该题满分，0~1）。
# 依据：教参通行的「难度系数」分档（P≥0.85 容易 / 0.4≤P<0.85 中等 / P<0.4 偏难），
# 取 0.40 作为「全班几乎都错」的线，是因为 40% 以下意味着多数学生连一半分都拿不到。
ITEM_ALMOST_ALL_WRONG = 0.40
ITEM_GIVE_AWAY = 0.85          # P≥0.85：白给分，讲评时可以快速带过
ITEM_HARD = 0.60               # P<0.60：需要重讲的分界线（比「几乎都错」宽一档）

# 区分度 D = 高分组平均得分率 − 低分组平均得分率（与 grade_reports.py 的 P/D 口径一致）。
# 高低分组各取 27%：这是测量学里的经典取值（Kelley 1939），分组太大区分度会被中间段稀释。
# 依据：D≥0.30 是能用的题；D<0.20 说明会的人和不会的人在这道题上没拉开差距，
# 要么题有歧义，要么这题考的东西全班都没掌握。
DISCRIM_GOOD = 0.30
DISCRIM_POOR = 0.20
DISCRIM_GROUP_FRAC = 0.27

# 连续下滑预警。
# 依据：总分满分 600，40 分≈6.7%，是一次月考里「掉出一个档位」的量级；
# 单次下滑 ≤ -8 分（≈1.3%）才算「这一次是下滑」，否则按正常起伏不记；
# 连续两次都下滑且累计掉 40 分以上才预警 —— 因为一次大波动不等于趋势。
DROP_MIN_STEP = 8.0
DROP_MIN_TOTAL = 40.0
DROP_CONSECUTIVE = 2           # 至少连续下滑几次

# 大起大落：任意一次上升或下降的幅度达到 40 分就记为「跳变」。
# 依据：与预警线同一量级（40 分），但方向要求「上下都出现过」——
# 只涨不跌不是大起大落，只跌不涨属于稳步下滑。
VOLATILE_SWING = 40.0

# 教学复盘：班级得分率与年级同知识点均值相差多少才值得说。
# 依据：5 个百分点是「一个班 40 人里差 2 人做对同一道题」量级，低于这个数的差
# 大概率是抽样波动（尤其知识点只挂一两道题时），说出来就是噪声。
GAP_MIN = 0.05

# 教学复盘的知识点层：这个点上至少 2 道题，得分率差才当作「讲法证据」。
# 依据：只挂 1 道题时，得分率就是那一道题的对错分布，题目本身的难度与歧义会直接
# 变成「讲法问题」；2 道以上才有「同一个点的不同问法」能互相印证。（比 MIN_ITEMS=3
# 宽松，是因为难点预判要给「第一道坎」这种更硬的结论，复盘只做对比。）
KP_MIN_ITEMS = 2

# 数据质检：单科 0 分但总分不为 0 时，只标记不排除（可能是缺考某科）。
QC_ZERO_TOTAL = 0.0            # 总分为 0 → 判缺考/未参加，排除出统计

# 禁语：教学复盘里绝不允许出现的「评价老师个人」的词。
# 依据：复盘的目的只有一个——下一节课怎么讲。一旦出现对人的评价，工具就变成监控系统。
FORBID_WORDS = [
    "教学水平", "教得好", "教得差", "不会教", "教不好", "能力差", "不负责", "失职",
    "态度不端", "敷衍", "绩效", "考核", "扣钱", "扣绩效", "该反省", "检讨",
    "你不行", "不合格", "误人子弟", "不称职",
]

SUBJECT_HINTS = ("语文", "数学", "英语", "物理", "化学", "道法", "历史", "地理", "生物")

# 各科满分：与 grade_reports.py 的 SUBJECTS 口径一致（语文/物理/化学 100，数学/英语 150）。
# 有这个表就用真满分算得分率；表里没有的科目退回「用年级最高分当满分」并在输出里标明。
FULL_BY_SUBJECT = {"语文": 100.0, "数学": 150.0, "英语": 150.0,
                   "物理": 100.0, "化学": 100.0}


def subject_full(h, all_scores):
    """科目满分：先查表，查不到用年级最高分兜底（兜底时调用方要在输出里说明）。"""
    for k, v in FULL_BY_SUBJECT.items():
        if k in h:
            return v, True
    top = max([x for x in all_scores if x is not None] or [0]) or 1.0
    return float(top), False


class InsightError(Exception):
    """可读的业务错误：命令行里直接打印 message，不带 traceback。"""


# ══════════════════════════════════════════════════════════════════════════
# 小工具
# ══════════════════════════════════════════════════════════════════════════
def num(v, default=None):
    """把单元格里的东西变成 float；空/None/非数字都回 default，不抛。"""
    if v is None:
        return default
    if isinstance(v, bool):
        return default
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("分", "")
    if not s:
        return default
    try:
        return float(s)
    except ValueError:
        return default


def pct(x, nd=0):
    return "—" if x is None else ("%.*f%%" % (nd, x * 100))


def mask(name):
    """名单里只留姓，其余打星：质检提示、控制台摘要用，避免整屏真名。"""
    s = (name or "").strip()
    if len(s) <= 1:
        return s or "（空）"
    return s[0] + "*" * (len(s) - 1)


def stamp():
    return datetime.datetime.now().strftime("%Y%m%d-%H%M%S")


def out_path(name):
    os.makedirs(OUT_DIR, exist_ok=True)
    return os.path.join(OUT_DIR, "%s_%s.xlsx" % (name, stamp()))


def guard_text(*texts):
    """禁语自检：教学复盘的所有文案在写出之前都要过这一关。

    命中就抛错——宁可让工具停下来，也不能让「评价老师个人」的话进报告。
    """
    blob = " ".join(str(t) for t in texts if t)
    hits = [w for w in FORBID_WORDS if w in blob]
    if hits:
        raise InsightError("输出文案里出现了评价老师个人的词：%s —— 复盘只谈知识点和讲法，"
                           "这几句要重写。" % "、".join(hits))
    return len(FORBID_WORDS)


def read_table(path):
    """读 csv/txt/xlsx 为 (head, rows)。csv 先试 utf-8-sig，失败再试 gbk。"""
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb.worksheets[0]
        rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        rows = [[("" if c is None else c) for c in r] for r in rows]
    elif ext in (".csv", ".txt"):
        text = None
        for enc in ("utf-8-sig", "utf-8", "gbk"):
            try:
                with open(path, "r", encoding=enc, newline="") as f:
                    text = f.read()
                break
            except UnicodeDecodeError:
                continue
        if text is None:
            raise InsightError("读不了这个文件（编码不是 utf-8 / gbk）：%s" % path)
        sample = text[:2048]
        delim = "\t" if sample.count("\t") > sample.count(",") else ","
        rows = [r for r in csv.reader(text.splitlines(), delimiter=delim)]
    else:
        raise InsightError("只认 .xlsx / .csv / .txt，这个是：%s" % path)
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise InsightError("文件是空的：%s" % path)
    head = [str(c).strip() for c in rows[0]]
    return head, rows[1:]


def col_of(head, *names):
    """按表头找列号（精确匹配优先，再退前后缀包含）。找不到返回 -1。"""
    for n in names:
        for i, h in enumerate(head):
            if h == n:
                return i
    for n in names:
        for i, h in enumerate(head):
            if h and (n in h):
                return i
    return -1


def is_subject_col(h):
    return any(s in h for s in SUBJECT_HINTS) and ("分数" in h or "成绩" in h)


# ══════════════════════════════════════════════════════════════════════════
# 题库
# ══════════════════════════════════════════════════════════════════════════
def load_bank(path=None):
    path = path or BANK
    if not os.path.exists(path):
        raise InsightError("找不到题库：%s" % path)
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as e:
        raise InsightError("题库读不动（%s）：%s" % (path, e))
    items = data.get("items") or []
    out = []
    for it in items:
        out.append({
            "no": str(it.get("no", "")).strip(),
            "chapter": str(it.get("chapter", "")).strip(),
            "kp": str(it.get("kp", "")).strip() or "未标注",
            "full": num(it.get("full"), 0.0) or 0.0,
            "difficulty": str(it.get("difficulty", "")).strip(),
            "err": str(it.get("err", "")).strip(),
            "stem": str(it.get("stem", "") or "").strip(),   # 可能是空串＝题面未录入
            "answer": str(it.get("answer", "") or "").strip(),
        })
    return {"meta": data.get("meta") or {}, "items": out}


def pick_items(bank, kp=None, chapter=None):
    items = bank["items"]
    if kp:
        return [i for i in items if i["kp"] == kp]
    if chapter:
        ch = str(chapter).strip()
        return [i for i in items if i["chapter"] == ch or i["chapter"].startswith(ch + ".")]
    return []


def stem_or_placeholder(it):
    """题面可能没录（题库里 7 条是空串）：统一成一句占位说明，不假装有题面。"""
    return it["stem"] if it["stem"] else "（题面未录入，只有题号与知识点）"


def no_label(no):
    s = str(no).strip()
    if s.startswith("第"):
        s = s[1:]
    if s.endswith("题"):
        s = s[:-1]
    return s


# ══════════════════════════════════════════════════════════════════════════
# 数据质检 + 读成绩单
# ══════════════════════════════════════════════════════════════════════════
def qc_report():
    return {"rows_in": 0, "rows_ok": 0, "excluded": [], "warn": [], "notes": []}


def load_grades(path):
    """读一份成绩单 -> {"path","label","students":[...],"qc":{...},"subjects":[...]}。

    质检口径：总分为 0 或缺 → 判缺考/未参加，排除出所有统计（否则均分被拉低）；
    单科 0 但总分正常 → 只提示不排除（可能是该科缺考，总分仍有效）。
    """
    if not os.path.exists(path):
        raise InsightError("找不到成绩单：%s" % path)
    # 读表走 ingest：多 sheet、表头容错（不只第一行）、合并单元格填满。
    # 只有「姓名 + 总分」都在的 sheet 才算成绩表；各 sheet 列布局必须一致，否则跳过并记一笔。
    rows_all, hdr, layout, notes = [], None, None, []
    for sheet in ingest.list_sheets(path):
        try:
            rows = ingest.sheet_rows(path, sheet)
        except Exception:
            continue
        hi, h = ingest.find_header(rows)
        if h is None:
            continue
        cmap = ingest.canon(h)
        if cmap.get("name") is None or cmap.get("total") is None:
            continue
        if layout is None:
            layout = cmap
            hdr = [("" if c is None else str(c).strip()) for c in h]
        elif cmap != layout:
            notes.append("跳过 sheet「%s」：列布局与前一个成绩表不同" % sheet)
            continue
        rows_all.extend(rows[hi + 1:])
    if layout is None:
        raise InsightError("这份表里没有同时含「姓名」「总分」表头的 sheet：%s" % path)

    def _col(key, *names):
        i = col_of(hdr, *names)
        if i is None or i < 0:
            i = layout.get(key, -1)
        return i if i is not None else -1

    c_name = _col("name", "姓名")
    c_cls = _col("cls", "班级")
    c_total = _col("total", "总分")          # 精确匹配优先：不会误取「小总分」
    subj_cols = [(h, i) for i, h in enumerate(hdr)
                 if i not in (c_name, c_cls, c_total) and is_subject_col(h)]

    qc = qc_report()
    qc["notes"].extend(notes)
    students, warn = [], []
    idx = 0
    for r in rows_all:
        idx += 1
        qc["rows_in"] += 1
        if not any(c is not None and str(c).strip() for c in r):
            continue
        name = str(r[c_name]).strip() if c_name >= 0 and c_name < len(r) and r[c_name] is not None else ""
        cls = str(r[c_cls]).strip() if c_cls >= 0 and c_cls < len(r) and r[c_cls] is not None else ""
        total = num(r[c_total]) if c_total >= 0 and c_total < len(r) else None
        subs = {}
        for h, i in subj_cols:
            subs[h] = num(r[i]) if i < len(r) else None
        if not name:
            qc["excluded"].append({"row": idx + 1, "name": "（空姓名）", "why": "没有姓名，无法归属"})
            continue
        if total is None:
            qc["excluded"].append({"row": idx + 1, "name": name, "why": "总分空着（缺考/未录分）"})
            continue
        if total <= QC_ZERO_TOTAL:
            qc["excluded"].append({"row": idx + 1, "name": name, "why": "总分 0（缺考/未参加）"})
            continue
        if name in {s["name"] for s in students}:
            warn.append("%s 在第 %d 行重复出现：按后一条算，前一条已丢弃" % (mask(name), idx + 1))
            students = [s for s in students if s["name"] != name]
        zero_sub = [h for h, v in subs.items() if v is not None and v <= 0]
        students.append({"name": name, "cls": cls, "total": total, "subs": subs,
                         "zero_sub": zero_sub})
    qc["rows_ok"] = len(students)
    qc["warn"] = warn
    qc["zero_sub_n"] = sum(1 for s in students if s["zero_sub"])
    label = os.path.splitext(os.path.basename(path))[0]
    return {"path": path, "label": label, "students": students, "qc": qc,
            "subjects": [h for h, _ in subj_cols]}


def load_scores(path, bank=None):
    """读逐题得分 -> 长表 [{student,no,kp,full,got,cls}]。

    认三种：
      1) 学生,题号,知识点,满分,得分   （grade_reports.py 的归档口径）
      2) 学生,题号,得分             （知识点/满分从题库补）
      3) 宽表：姓名/学生 + 每列一个题号（值必须是数字分）
    """
    bank = bank or load_bank()
    full_of = {it["no"]: it["full"] for it in bank["items"]}
    kp_of = {it["no"]: it["kp"] for it in bank["items"]}
    head, rows = read_table(path)
    qc = qc_report()
    qc["rows_in"] = len(rows)
    out, warn = [], []
    c_stu = col_of(head, "学生", "姓名", "名字")
    c_no = col_of(head, "题号", "题目", "编号")
    c_got = col_of(head, "得分", "分数", "成绩")
    c_full = col_of(head, "满分", "分值")
    c_kp = col_of(head, "知识点", "考点")
    c_cls = col_of(head, "班级", "班")
    wide = c_stu >= 0 and c_no < 0 and c_got < 0
    if c_stu < 0:
        raise InsightError("逐题得分表里找不到「学生」或「姓名」列：%s" % path)

    if wide:
        # 宽表：除身份列外，每列当一道题
        qnos = [h for i, h in enumerate(head) if i not in (c_stu, c_cls) and h]
        for r in rows:
            stu = str(r[c_stu]).strip() if c_stu < len(r) and r[c_stu] is not None else ""
            if not stu:
                continue
            cls = str(r[c_cls]).strip() if 0 <= c_cls < len(r) and r[c_cls] is not None else ""
            for i, no in enumerate(head):
                if i in (c_stu, c_cls) or not no:
                    continue
                v = r[i] if i < len(r) else None
                got = num(v)
                if got is None:
                    if str(v or "").strip():
                        warn.append("题 %s 的单元格不是数字（%r），按未作答记 0 分" % (no, v))
                    got = 0.0
                full = full_of.get(no_label(no), 0.0) or 1.0
                if got > full:
                    warn.append("%s 题 %s 得分 %g 超过满分 %g：按满分封顶" % (mask(stu), no, got, full))
                    got = full
                out.append({"student": stu, "cls": cls, "no": no_label(no),
                            "kp": kp_of.get(no_label(no), "未标注"), "full": full, "got": got})
    else:
        if c_got < 0:
            raise InsightError("逐题得分表里找不到「得分」列：%s" % path)
        for r in rows:
            stu = str(r[c_stu]).strip() if c_stu < len(r) and r[c_stu] is not None else ""
            no = no_label(r[c_no]) if 0 <= c_no < len(r) and r[c_no] is not None else ""
            if not stu or not no:
                continue
            got = num(r[c_got]) if c_got < len(r) else None
            if got is None:
                got = 0.0
            full = num(r[c_full]) if 0 <= c_full < len(r) else None
            if not full:
                full = full_of.get(no, 0.0)
            if not full:
                warn.append("题 %s 在题库里查不到满分，按 1 分算" % no)
                full = 1.0
            if got > full:
                warn.append("%s 题 %s 得分 %g 超过满分 %g：按满分封顶" % (mask(stu), no, got, full))
                got = full
            kp = str(r[c_kp]).strip() if 0 <= c_kp < len(r) and r[c_kp] not in (None, "") else ""
            cls = str(r[c_cls]).strip() if 0 <= c_cls < len(r) and r[c_cls] is not None else ""
            out.append({"student": stu, "cls": cls, "no": no,
                        "kp": kp or kp_of.get(no, "未标注"), "full": full, "got": got})
    qc["rows_ok"] = len(out)
    qc["warn"] = warn[:40]
    qc["students"] = len({r["student"] for r in out})
    qc["items"] = len({r["no"] for r in out})
    return {"path": path, "label": os.path.splitext(os.path.basename(path))[0],
            "rows": out, "qc": qc, "mode": "宽表" if wide else "长表"}


def attach_class(scores, grades):
    """逐题表里没有班级时，用成绩单按姓名补上（改名/重名不猜，补不上的留空）。"""
    if not grades:
        return 0
    by_name = {}
    for s in grades["students"]:
        by_name.setdefault(s["name"], s["cls"])
    filled = 0
    for r in scores["rows"]:
        if not r.get("cls"):
            c = by_name.get(r["student"])
            if c:
                r["cls"] = c
                filled += 1
    return filled


# ══════════════════════════════════════════════════════════════════════════
# 逐题统计
# ══════════════════════════════════════════════════════════════════════════
def item_table(rows):
    """按题号聚合：n / 平均得分 / 得分率 P / 区分度 D / 高分低分组的得分率。"""
    by_no = {}
    per_stu = {}
    for r in rows:
        b = by_no.setdefault(r["no"], {"no": r["no"], "kp": r["kp"], "full": r["full"],
                                       "got": [], "got_by_stu": {}})
        b["got"].append(r["got"])
        b["full"] = max(b["full"], r["full"]) if b["full"] else r["full"]
        b["got_by_stu"][r["student"]] = r["got"]
        p = per_stu.setdefault(r["student"], {"got": 0.0, "full": 0.0})
        p["got"] += r["got"]
        p["full"] += r["full"]
    rates = {k: (v["got"] / v["full"] if v["full"] else 0.0) for k, v in per_stu.items()}
    order = sorted(rates, key=lambda k: (-rates[k], k))
    n = len(order)
    k = max(1, int(round(n * DISCRIM_GROUP_FRAC))) if n >= 3 else 1
    top, bot = set(order[:k]), set(order[-k:])
    out = []
    for no, b in sorted(by_no.items(), key=lambda kv: str(kv[0])):
        n_b = len(b["got"])
        avg = sum(b["got"]) / n_b if n_b else 0.0
        p_val = avg / b["full"] if b["full"] else 0.0
        hi = [b["got_by_stu"][s] for s in top if s in b["got_by_stu"]]
        lo = [b["got_by_stu"][s] for s in bot if s in b["got_by_stu"]]
        if b["full"] and hi and lo and len(rates) >= 3:
            d = (sum(hi) / len(hi) - sum(lo) / len(lo)) / b["full"]
        else:
            d = None
        out.append({"no": no, "kp": b["kp"], "full": b["full"], "n": n_b,
                    "avg": avg, "P": p_val, "D": d,
                    "n_right": sum(1 for g in b["got"] if b["full"] and g >= b["full"] - 1e-9),
                    "n_zero": sum(1 for g in b["got"] if g <= 0)})
    return out


def judge_base(it):
    """只看得分率的那半句判定（预判句里用它，别把区分度的括号塞进人话里）。"""
    if it["P"] <= ITEM_ALMOST_ALL_WRONG:
        return "全班几乎都错"
    if it["P"] < ITEM_HARD:
        return "偏难，要重讲"
    if it["P"] >= ITEM_GIVE_AWAY:
        return "白给分"
    return "正常"


def judge_item(it):
    """完整判定：得分率 + 区分度。区分度缺失（学生太少）时不给区分度结论。"""
    base = judge_base(it)
    if it["D"] is None:
        return base + "（人数太少，区分度算不出）"
    if it["D"] < DISCRIM_POOR:
        return base + "｜区分度低"
    return base


def kp_table(rows):
    """按知识点聚合：满分合计 / 得分合计 / 得分率 / 学生数。"""
    acc = {}
    for r in rows:
        b = acc.setdefault(r["kp"], {"kp": r["kp"], "full": 0.0, "got": 0.0,
                                     "nos": set(), "stus": set()})
        b["full"] += r["full"]
        b["got"] += r["got"]
        b["nos"].add(r["no"])
        b["stus"].add(r["student"])
    out = []
    for b in acc.values():
        out.append({"kp": b["kp"], "full": b["full"], "got": b["got"],
                    "P": b["got"] / b["full"] if b["full"] else 0.0,
                    "n_items": len(b["nos"]), "n_students": len(b["stus"]),
                    "nos": sorted(b["nos"], key=lambda x: str(x))})
    return sorted(out, key=lambda x: x["P"])


# ══════════════════════════════════════════════════════════════════════════
# 知识点 → 前置 / 讲法模板（本地知识，不花模型）
# ══════════════════════════════════════════════════════════════════════════
PRE_REQ = {
    "绝对值": ["能说出「一个数的绝对值＝它在数轴上到原点的距离」，而不是背「去掉负号」",
               "会在数轴上标出 a 与 -a，并说出两者到原点的距离相等"],
    "数轴": ["能分清「一个点的位置」和「两点之间的距离」是两件事",
             "会读「原点、正方向、单位长度」三要素，能在轴上标负数"],
    "相反数": ["知道相反数是关于原点对称的位置关系，不是「长得像」",
               "会把「-(-5)」读成「-5 的相反数」而不是连减两个负号"],
    "有理数的概念与分类": ["能按「整数/分数、正/负/零」两个维度给同一个数定位",
                           "知道 0 既不是正数也不是负数，分类表要先立住"],
    "有理数的大小比较": ["会用数轴比大小（右大左小），再过渡到法则",
                         "会处理负数比较：绝对值大的反而小"],
    "有理数加法": ["先过「同号相加 / 异号相加」的符号判定，再谈绝对值加减",
                   "会用数轴走格子解释「-7 + 3」为什么是 -4"],
    "有理数减法": ["知道「减去一个数＝加上它的相反数」，并且两项一起变号",
                   "能在一步里同时改减号和后面那个数的符号，不掉一个"],
    "有理数乘除": ["先确认符号法则「负负得正」，再谈绝对值相乘除",
                   "会把除法写成乘倒数，避免除号连写"],
    "有理数乘方": ["知道 (-2)³ 与 -2³ 的区别（底数带不带括号）",
                   "会数「负数的奇次幂为负、偶次幂为正」"],
    "有理数混合运算": ["先立运算顺序：乘方 > 乘除 > 加减，同级从左到右",
                       "会先把每一步的符号定下来，再算绝对值"],
    "有理数的实际应用": ["会把「上升/下降、零上/零下」翻译成正负数再列式",
                         "会检查答案的量纲与生活常识（温度不会算成 -80℃）"],
    "正负数概念": ["知道正负号在这里表示「方向/意义相反」，不只是运算符号",
                   "能举出生活里的正负例子（收支、海拔、温度）"],
    "倒数": ["知道倒数是「相乘等于 1」的那对数，不是相反数",
             "会处理负倒数：-1/3 的倒数是 -3"],
    "科学记数法": ["会数小数点挪了几位，再决定 10 的指数",
                   "知道 a×10ⁿ 里 1≤|a|<10，指数与位数差一"],
    "近似数": ["知道精确到哪一位要看那一位的下一位",
               "会写有效数字，能区分「精确到千分位」和「保留三位小数」"],
}


def pre_req_of(kp, chapter=""):
    """知识点的前置确认项；没收录的按章节给通用前置，绝不硬编。"""
    if kp in PRE_REQ:
        return PRE_REQ[kp]
    if chapter:
        return ["能用自己的话说出「%s」这节要解决什么问题" % kp,
                "会做本节里最简单的那个例题（不看答案，独立写一遍）"]
    return ["能用自己的话说出「%s」的定义" % kp,
            "会做「%s」的最基础那一道题，不看答案" % kp]


ERR_CAN = {
    "概念不清": "坎不在算，在「这个点到底是什么意思」——定义没落地，练多少道都是照猫画虎",
    "步骤遗漏": "坎在「心里有、纸上没」——步骤一多就跳环，得分点丢在半路",
    "计算错误": "坎在「会做但算错」——符号、进位这类地方反复失手，方法本身不是问题",
    "审题错误": "坎在「读题就跑」——条件和问题对不上，一动手就偏",
}


def err_mix(items):
    """题库里这个知识点挂的错因倾向分布（命题人预设，不是实测）。"""
    acc = {}
    for it in items:
        e = it["err"] or "未标注"
        acc[e] = acc.get(e, 0) + 1
    return sorted(acc.items(), key=lambda kv: (-kv[1], kv[0]))


# ══════════════════════════════════════════════════════════════════════════
# xlsx 输出
# ══════════════════════════════════════════════════════════════════════════
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def write_xlsx(path, sheets):
    """sheets = [(sheet_name, head_list, rows_list_of_list, note_or_None)]"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, head, rows, note in sheets:
        ws = wb.create_sheet(name[:31])
        r0 = 1
        if note:
            ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="808080")
            ws.merge_cells(start_row=1, start_column=1,
                           end_row=1, end_column=max(1, len(head)))
            r0 = 2
        for j, h in enumerate(head, 1):
            c = ws.cell(row=r0, column=j, value=h)
            c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, row in enumerate(rows, r0 + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(row=i, column=j, value=v)
                c.border = BORDER
                if isinstance(v, str) and len(v) > 24:
                    c.alignment = Alignment(wrap_text=True, vertical="top")
        for j, h in enumerate(head, 1):
            width = max(len(str(h)) * 2 + 4, 10)
            for row in rows[:80]:
                if j <= len(row):
                    width = max(width, min(60, len(str(row[j - 1])) * 2 + 2))
            ws.column_dimensions[get_column_letter(j)].width = min(60, width)
        ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return path


def print_qc(qc, label):
    """数据质检的统一打印。排除项一律点名到位数，别让均分悄悄被拉低。"""
    if not qc:
        return
    print("  数据质检〔%s〕：读入 %d 行 → 参与统计 %d 行" % (
        label, qc.get("rows_in", 0), qc.get("rows_ok", 0)))
    ex = qc.get("excluded") or []
    if ex:
        head = "、".join("%s(第%d行,%s)" % (mask(e["name"]), e["row"], e["why"]) for e in ex[:6])
        more = "" if len(ex) <= 6 else " 等 %d 人" % len(ex)
        print("    已排除 %d 行：%s%s" % (len(ex), head, more))
    if qc.get("zero_sub_n"):
        print("    另有 %d 人单科为 0（总分有效，只提示不排除）" % qc["zero_sub_n"])
    for w in (qc.get("warn") or [])[:6]:
        print("    提示：%s" % w)


# ══════════════════════════════════════════════════════════════════════════
# 命令 1：难点预判
# ══════════════════════════════════════════════════════════════════════════
def cmd_hardkp(args):
    bank = load_bank(args.bank)
    if not args.kp and not args.chapter:
        raise InsightError("要给 --kp 或 --chapter，例如 --kp 绝对值")
    items = pick_items(bank, kp=args.kp, chapter=args.chapter)
    title = args.kp or ("第 %s 章" % args.chapter)
    print("=" * 78)
    print("难点预判：%s" % title)
    print("=" * 78)
    if not items:
        kps = sorted({i["kp"] for i in bank["items"]})
        raise InsightError("题库里没有「%s」。现有知识点：%s" % (title, "、".join(kps)))

    print("  题库里挂在这个点上的题：%d 道（满分合计 %g 分）" % (
        len(items), sum(i["full"] for i in items)))
    for it in items:
        print("    第%s题 %-6s %g分 难度%-4s 错因倾向 %-6s %s" % (
            it["no"], it["chapter"], it["full"], it["difficulty"] or "-",
            it["err"] or "-", (it["stem"] or "（题面未录入）")[:24]))

    short = len(items) < MIN_ITEMS
    if short:
        print()
        print("  【样本不足】这个点只有 %d 道题，低于下限 %d 道——不做难点判定、不给预判。" % (
            len(items), MIN_ITEMS))
        print("  为什么：%d 道题分不清「这个点本身就难」还是「这一道题出偏了」；" % len(items))
        print("  补到 %d 道（同一知识点的不同问法）再跑一次，结论才站得住。" % MIN_ITEMS)

    # 逐题得分（可选）
    stats, sres = {}, None
    if args.scores:
        sres = load_scores(args.scores, bank)
        print()
        print_qc(sres["qc"], "逐题得分 %s，%s" % (sres["label"], sres["mode"]))
        nos = {it["no"] for it in items}
        sub = [r for r in sres["rows"] if r["no"] in nos]
        if not sub:
            print("    这份逐题得分里没有这个点的题（题号对不上题库），判定跳过。")
        else:
            stats = {t["no"]: t for t in item_table(sub)}
    else:
        print()
        print("  【没有逐题得分】没给 --scores，得分率与区分度算不出来，下面是题库侧的判断。")

    rows_out = []
    worst = None
    for it in items:
        t = stats.get(it["no"])
        if t:
            rows_out.append([it["no"], it["kp"], it["chapter"], it["full"], t["n"],
                             round(t["avg"], 1), pct(t["P"], 1),
                             "—" if t["D"] is None else round(t["D"], 2),
                             judge_item(t), it["difficulty"] or "", it["err"] or "",
                             stem_or_placeholder(it)[:40]])
            if worst is None or t["P"] < worst[1]["P"]:
                worst = (it, t)
        else:
            rows_out.append([it["no"], it["kp"], it["chapter"], it["full"], 0, "—", "—", "—",
                             "无得分数据", it["difficulty"] or "", it["err"] or "",
                             stem_or_placeholder(it)[:40]])

    if stats:
        print()
        print("  %-4s %-6s %-5s %-5s %-7s %-6s %-8s %s" % (
            "题号", "满分", "人数", "平均分", "得分率", "区分度", "判定", "题面"))
        for it in items:
            t = stats.get(it["no"])
            if not t:
                print("  第%s题 %-5g %-5s %-5s %-7s %-6s %-8s %s" % (
                    it["no"], it["full"], "—", "—", "—", "—", "无得分数据",
                    (it["stem"] or "（题面未录入）")[:20]))
                continue
            print("  第%s题 %-5g %-5d %-5.1f %-7s %-6s %-8s %s" % (
                it["no"], it["full"], t["n"], t["avg"], pct(t["P"], 1),
                "—" if t["D"] is None else "%.2f" % t["D"], judge_item(t),
                (it["stem"] or "（题面未录入）")[:20]))

    # 预判
    print()
    mix = err_mix(items)
    top_err = mix[0][0] if mix else "未标注"
    verdict_lines, pre_lines = [], []
    if short:
        verdict = ("【不下结论】%s在题库里只有 %d 道题（< %d 道），还看不出第一道坎在哪；"
                   "先按分数把题补齐。" % (title, len(items), MIN_ITEMS))
        verdict_lines.append(verdict)
        print("  一句话预判：%s" % verdict)
        print("  必须先确认的两个前置：暂不给——题量不够，给了也是猜。")
    else:
        if stats and worst:
            it_w, t_w = worst
            hard_n = sum(1 for t in stats.values() if t["P"] <= ITEM_ALMOST_ALL_WRONG)
            verdict = ("%s的第一道坎是「%s」：本点 %d 道题里 %d 道得分率 ≤ %s，最低的第 %s 题只有 %s"
                       "（%s）——%s。" % (
                           title, it_w["err"] or top_err, len(items), hard_n,
                           pct(ITEM_ALMOST_ALL_WRONG), it_w["no"], pct(t_w["P"], 1),
                           judge_base(t_w), ERR_CAN.get(it_w["err"] or top_err, ERR_CAN["概念不清"])))
        elif len(mix) > 1 and mix[0][1] == mix[1][1]:
            verdict = ("%s的题库侧错因是分散的（%s 各 %d 道），没有单一主导错因："
                       "这种情况先不按错因下结论，把每个错因各挑一道最基础的题验一遍，"
                       "看学生实际卡在哪；缺逐题得分，这里只到命题预设为止。" % (
                           title, "、".join(e for e, n in mix if n == mix[0][1]), mix[0][1]))
        else:
            verdict = ("%s的题库侧错因集中在「%s」（%d / %d 道）：%s。"
                       "缺逐题得分，这一句是命题预设，不是这次考试的实测。" % (
                           title, top_err, mix[0][1], len(items),
                           ERR_CAN.get(top_err, "先按定义落地再练")))
        verdict_lines.append(verdict)
        print("  一句话预判：%s" % verdict)
        pre = []
        for it in items:                    # 章节模式下这一章可能有多个知识点，先去重
            for p in pre_req_of(it["kp"], it["chapter"]):
                if p not in pre:
                    pre.append(p)
        pre = pre[:2]
        pre_lines = pre
        print("  必须先确认的两个前置：")
        for i, p in enumerate(pre, 1):
            print("    %d) %s" % (i, p))
        # 讲法建议（只讲知识点，不评价人）
        if stats:
            give = [it for it in items if stats.get(it["no"]) and stats[it["no"]]["P"] >= ITEM_GIVE_AWAY]
            hard = [it for it in items if stats.get(it["no"]) and stats[it["no"]]["P"] <= ITEM_ALMOST_ALL_WRONG]
            if give and hard:
                print("  讲评次序：第 %s 题（%s）10 秒带过；把时间留给第 %s 题（%s）——同一个点，"
                      "一个会一个不会，差别就在那一步。" % (
                          "、".join(g["no"] for g in give), pct(ITEM_GIVE_AWAY),
                          "、".join(h["no"] for h in hard), pct(ITEM_ALMOST_ALL_WRONG)))
            elif hard:
                print("  讲评次序：这个点没有白给分的题，第 %s 题全班几乎都错（≤ %s）——"
                      "下节课整块时间就给这一个点，别急着往后赶。" % (
                          "、".join(h["no"] for h in hard), pct(ITEM_ALMOST_ALL_WRONG)))
            elif give and len(items) > len(give):
                print("  讲评次序：第 %s 题已是白给分（≥ %s），这个点只需快速核对；"
                      "把课时留给下一个点。" % (
                          "、".join(g["no"] for g in give), pct(ITEM_GIVE_AWAY)))
        good_d = [it for it in items if stats.get(it["no"]) and stats[it["no"]]["D"] is not None
                  and stats[it["no"]]["D"] >= DISCRIM_GOOD]
        if good_d:
            print("  可作分层题源：第 %s 题区分度 ≥ %.1f（会的人做对、不会的人做错）——"
                  "抽出来放提高组，别全班重复讲。" % (
                      "、".join(x["no"] for x in good_d), DISCRIM_GOOD))
        low_d = [it for it in items if stats.get(it["no"]) and stats[it["no"]]["D"] is not None
                 and stats[it["no"]]["D"] < DISCRIM_POOR]
        if low_d:
            print("  注意命题：第 %s 题区分度 < %.1f，会的人和不会的人在这道题上没拉开——"
                  "要么题有歧义，要么这一步全班都没真的掌握。" % (
                      "、".join(x["no"] for x in low_d), DISCRIM_POOR))

    note = "数据来源：题库.json + %s。阈值口径见 insight.py 顶部常量。" % (
        os.path.basename(args.scores) if args.scores else "（未提供逐题得分）")
    path = out_path("难点预判_%s" % (args.kp or ("第%s章" % args.chapter)))
    write_xlsx(path, [
        ("题目判定", ["题号", "知识点", "章节", "满分", "作答人数", "平均分", "得分率",
                     "区分度D", "判定", "预设难度", "错因倾向", "题面"],
         rows_out, note),
        ("预判与前置", ["项", "内容"],
         [["一句话预判", verdict_lines[0]]] +
         [["前置确认 %d" % i, p] for i, p in enumerate(pre_lines, 1)] +
         [["数据质检", "读入 %d 行 → 参与统计 %d 行；排除 %d 行" % (
             (sres or {}).get("qc", {}).get("rows_in", 0),
             (sres or {}).get("qc", {}).get("rows_ok", 0),
             len((sres or {}).get("qc", {}).get("excluded", [])))],
          ["口径", note],
          ["规矩", "样本不足就明说不足；阈值一律写在顶部常量里，代码里不出现魔法数字"]],
         None),
    ])
    print()
    print("  → 已落表：%s" % path)
    return 0


# ══════════════════════════════════════════════════════════════════════════
# 命令 2：连续下滑预警
# ══════════════════════════════════════════════════════════════════════════
def cmd_drop(args):
    files = list(args.files or [])
    if not files:
        raise InsightError("要给至少一份成绩单：--files 第一次.xlsx 第二次.xlsx")
    for f in files:
        if not os.path.exists(f):
            raise InsightError("找不到成绩单：%s" % f)
    print("=" * 78)
    print("连续下滑预警：共 %d 次考试" % len(files))
    print("=" * 78)
    exams = []
    for f in files:
        g = load_grades(f)
        exams.append(g)
        print_qc(g["qc"], os.path.basename(f))

    labels = [e["label"] for e in exams]
    print("  场次：%s" % " → ".join(labels))

    if len(exams) == 1:
        print()
        print("  【趋势算不出来】只有 1 次数据。相邻两次的总分变化至少要 2 份成绩单才算得出，")
        print("  这次只能看质检结果，不给下滑名单——不编。")
        path = out_path("下滑预警_数据不足")
        write_xlsx(path, [
            ("说明", ["项", "内容"],
             [["结论", "只有 1 次考试，趋势算不出来（相邻两次变化需要 ≥2 份成绩单）"],
              ["本次读入", "%s：%d 行 → 有效 %d 人" % (
                  labels[0], exams[0]["qc"]["rows_in"], exams[0]["qc"]["rows_ok"])],
              ["判定阈值", "连续 %d 次单次下滑 ≥ %.0f 分且累计 ≥ %.0f 分 → 稳步下滑；"
                          "上下都出现 ≥ %.0f 分的跳变 → 大起大落（见顶部常量注释）" % (
                              DROP_CONSECUTIVE, DROP_MIN_STEP, DROP_MIN_TOTAL, VOLATILE_SWING)]],
             "数据质检已执行：总分为 0 的行按缺考排除。"),
            ("数据质检", ["场次", "读入行", "有效行", "排除行", "排除名单"],
             [[e["label"], e["qc"]["rows_in"], e["qc"]["rows_ok"],
               len(e["qc"]["excluded"]),
               "、".join("%s(%s)" % (x["name"], x["why"]) for x in e["qc"]["excluded"])]
              for e in exams], None),
        ])
        print("  → 已落表：%s" % path)
        return 0

    # 合并：每人一条轨迹；某次缺考（被质检排除）→ 这条轨迹缺口，不给趋势
    order, seen = [], {}
    for ei, e in enumerate(exams):
        for s in e["students"]:
            if s["name"] not in seen:
                seen[s["name"]] = {"name": s["name"], "cls": s["cls"],
                                   "vals": [None] * len(exams), "subs": [None] * len(exams)}
                order.append(s["name"])
            seen[s["name"]]["vals"][ei] = s["total"]
            seen[s["name"]]["subs"][ei] = s["subs"]
            if not seen[s["name"]]["cls"]:
                seen[s["name"]]["cls"] = s["cls"]

    subjects = []
    for e in exams:
        for h in e["subjects"]:
            if h not in subjects:
                subjects.append(h)

    rows, gaps, steady, volatile = [], [], [], []
    for name in order:
        p = seen[name]
        vals = p["vals"]
        gaps_i = [i for i, v in enumerate(vals) if v is None]
        if gaps_i:
            gaps.append({"name": name, "cls": p["cls"], "vals": vals,
                         "why": "第 %s 次没有有效总分（缺考/0 分）" % (
                             "、".join(str(i + 1) for i in gaps_i))})
            continue
        deltas = [vals[i + 1] - vals[i] for i in range(len(vals) - 1)]
        tail = []
        for d in reversed(deltas):          # 从最后一次往前数「连续下滑」
            if d <= -DROP_MIN_STEP:
                tail.append(d)
            else:
                break
        steady_run = len(tail)
        steady_sum = sum(tail)
        rises = [d for d in deltas if d >= VOLATILE_SWING]
        falls = [d for d in deltas if d <= -VOLATILE_SWING]
        is_volatile = bool(rises and falls)
        is_steady = steady_run >= DROP_CONSECUTIVE and steady_sum <= -DROP_MIN_TOTAL
        # 下滑最大的那一科（首末对比，只用两次都有分的科目）
        subj_delta = {}
        for h in subjects:
            a = (p["subs"][0] or {}).get(h)
            b = (p["subs"][-1] or {}).get(h)
            if a is not None and b is not None:
                subj_delta[h] = b - a
        worst_sub = min(subj_delta.items(), key=lambda kv: kv[1]) if subj_delta else None
        row = {"name": name, "cls": p["cls"], "vals": vals,
               "total_change": vals[-1] - vals[0],
               "steady_run": steady_run, "steady_sum": steady_sum,
               "volatile": is_volatile, "steady": is_steady,
               "rises": rises, "falls": falls,
               "worst_sub": worst_sub, "subj_delta": subj_delta}
        rows.append(row)
        if is_volatile:
            volatile.append(row)
        elif is_steady:
            steady.append(row)

    print()
    print("  判定口径：单次下滑 ≥ %.0f 分才计入；连续 %d 次都下滑且累计 ≥ %.0f 分 → 稳步下滑；"
          % (DROP_MIN_STEP, DROP_CONSECUTIVE, DROP_MIN_TOTAL))
    print("            上下都出现过 ≥ %.0f 分的跳变 → 大起大落（两类干预方式不同，分开列）"
          % VOLATILE_SWING)
    print("  参与判定：%d 人；因缺考/0 分断档不判趋势：%d 人" % (len(rows), len(gaps)))

    def sub_brief(r):
        if not r["worst_sub"]:
            return "—"
        k, v = r["worst_sub"]
        if v < 0:
            return "最大下滑科 %s %+.0f" % (k, v)
        return "无下滑科目（最小的 %s 也是 %+.0f）" % (k, v)

    def brief(r):
        return "%s（%s）｜%s｜累计 %+.1f｜%s" % (
            r["name"], r["cls"] or "—", " / ".join("%.0f" % v for v in r["vals"]),
            r["total_change"], sub_brief(r))

    print()
    if steady:
        print("  【稳步下滑】%d 人 —— 不是一次失手，是慢慢滑。干预：查最近两周的落笔与错题重做，" % len(steady))
        print("              别先加题量；这类学生的问题通常在「每天都在漏一点」。")
        for r in sorted(steady, key=lambda x: x["steady_sum"])[:12]:
            print("    - %s｜连续下滑 %d 次合计 %+.1f" % (brief(r), r["steady_run"], r["steady_sum"]))
    else:
        print("  【稳步下滑】0 人 —— 没人满足「连续 %d 次下滑且累计 ≥ %.0f 分」。" % (
            DROP_CONSECUTIVE, DROP_MIN_TOTAL))
    if volatile:
        print()
        print("  【大起大落】%d 人 —— 不是趋势，是状态。干预：别按趋势做长期判断，" % len(volatile))
        print("              先看考试当天状态、某科是否临时放弃；这类波动靠盯「稳」不靠加量。")
        for r in sorted(volatile, key=lambda x: -(max(x["rises"]) - min(x["falls"])))[:12]:
            print("    - %s｜最大反弹 %+.0f、最大跌落 %+.0f" % (
                brief(r), max(r["rises"]), min(r["falls"])))
    else:
        print()
        print("  【大起大落】0 人 —— 没有人在一次上升 %.0f 分以上的同时又一次下跌 %.0f 分以上。" % (
            VOLATILE_SWING, VOLATILE_SWING))
    if gaps:
        print()
        print("  【数据断档，不判趋势】%d 人（缺考/0 分先排除，不能拿 0 分当「暴跌」）" % len(gaps))
        for g in gaps[:8]:
            print("    - %s（%s）｜%s" % (mask(g["name"]), g["cls"] or "—", g["why"]))

    note = ("判定阈值：单次下滑 %.0f 分 / 累计 %.0f 分 / 连续 %d 次 / 跳变量级 %.0f 分"
            "（依据见 insight.py 顶部常量注释）" % (
                DROP_MIN_STEP, DROP_MIN_TOTAL, DROP_CONSECUTIVE, VOLATILE_SWING))
    path = out_path("下滑预警_%d次" % len(exams))
    head = ["类型", "姓名", "班级"] + ["第%d次 %s" % (i + 1, l) for i, l in enumerate(labels)] + \
           ["累计变化", "连续下滑次数", "连续下滑合计", "下滑最大的科", "科目首末变化", "建议"]
    out_rows = []
    for r in sorted(steady, key=lambda x: x["steady_sum"]):
        out_rows.append(["稳步下滑", r["name"], r["cls"]] + [round(v, 1) for v in r["vals"]] +
                        [round(r["total_change"], 1), r["steady_run"], round(r["steady_sum"], 1),
                         (r["worst_sub"][0] if r["worst_sub"] and r["worst_sub"][1] < 0
                          else "（无科目下滑）"),
                         "、".join("%s%+.0f" % (k, v) for k, v in r["subj_delta"].items()),
                         "不是一次失手：查最近两周作业落笔与错题重做，先别加题量"])
    for r in sorted(volatile, key=lambda x: -(max(x["rises"]) - min(x["falls"]))):
        out_rows.append(["大起大落", r["name"], r["cls"]] + [round(v, 1) for v in r["vals"]] +
                        [round(r["total_change"], 1), r["steady_run"], round(r["steady_sum"], 1),
                         (r["worst_sub"][0] if r["worst_sub"] and r["worst_sub"][1] < 0
                          else "（无科目下滑）"),
                         "、".join("%s%+.0f" % (k, v) for k, v in r["subj_delta"].items()),
                         "看状态不看趋势：先查考试当天状态与该科是否临时放弃"])
    for g in gaps:
        out_rows.append(["数据断档", g["name"], g["cls"]] +
                        ["" if v is None else round(v, 1) for v in g["vals"]] +
                        ["", "", "", "", "", g["why"]])
    write_xlsx(path, [
        ("预警名单", head, out_rows, note),
        ("数据质检", ["场次", "读入行", "有效行", "排除行", "排除名单"],
         [[e["label"], e["qc"]["rows_in"], e["qc"]["rows_ok"], len(e["qc"]["excluded"]),
           "、".join("%s(%s)" % (x["name"], x["why"]) for x in e["qc"]["excluded"])]
          for e in exams], None),
    ])
    print()
    print("  → 已落表：%s" % path)
    return 0


# ══════════════════════════════════════════════════════════════════════════
# 命令 3：教学复盘
# ══════════════════════════════════════════════════════════════════════════
def cmd_reflect(args):
    files = list(args.files or [])
    if not files:
        raise InsightError("要给成绩单：--files 成绩单.xlsx")
    grades = load_grades(files[-1])          # 复盘看最近一次
    if len(files) > 1:
        print("  给了 %d 份成绩单，复盘只看最近一次：%s" % (len(files), grades["label"]))
    print("=" * 78)
    print("教学复盘：%s（只谈知识点与讲法，不做任何对人的评价）" % grades["label"])
    print("=" * 78)
    print_qc(grades["qc"], grades["label"])

    classes = {}
    for s in grades["students"]:
        if s["cls"]:
            classes.setdefault(s["cls"], []).append(s)
    if args.cls:
        if args.cls not in classes:
            raise InsightError("成绩单里没有这个班：%s。现有：%s" % (
                args.cls, "、".join(sorted(classes)) or "（没有班级列）"))
        unit = args.cls
        unit_note = "按 --class 指定"
    elif classes:
        unit = max(classes.items(), key=lambda kv: len(kv[1]))[0]
        unit_note = "未给 --class，默认取人数最多的班（可以 --class 换）"
    else:
        unit = ""
        unit_note = "成绩单没有班级列，全表当同一个单位看"
    unit_rows = [s for s in grades["students"] if (s["cls"] == unit if unit else True)]
    all_rows = grades["students"]
    print("  分析单位：%s（%d 人）｜年级有效样本 %d 人｜%s" % (
        unit or "全表", len(unit_rows), len(all_rows), unit_note))

    # ---- 科目层（真实成绩单就能算）----
    subj_stats, gaps, dropped_sub = [], [], []
    for h in grades["subjects"]:
        # 单科 0 分按缺考处理：不计入该科均分（否则一个缺考能把全班均分拉下去）
        u_all = [s["subs"][h] for s in unit_rows if s["subs"].get(h) is not None]
        a_all = [s["subs"][h] for s in all_rows if s["subs"].get(h) is not None]
        u = [x for x in u_all if x > 0]
        a = [x for x in a_all if x > 0]
        if len(u) < len(u_all):
            dropped_sub.append("%s %d 人" % (h, len(u_all) - len(u)))
        full, from_table = subject_full(h, a_all)
        if u and a:
            u_avg, a_avg = sum(u) / len(u), sum(a) / len(a)
            rate_u = u_avg / full if full else 0.0
            rate_a = a_avg / full if full else 0.0
            subj_stats.append({"h": h, "u_avg": u_avg, "a_avg": a_avg,
                               "rate_u": rate_u, "rate_a": rate_a,
                               "gap": rate_u - rate_a, "n": len(u), "full": full,
                               "full_from_table": from_table,
                               "share": round(len(u) * 100.0 / len(a), 1) if a else 0.0})
            if rate_u - rate_a <= -GAP_MIN:
                gaps.append({"kind": "科目", "name": h, "rate_u": rate_u, "rate_a": rate_a,
                             "gap": rate_u - rate_a, "n": len(u)})
    if dropped_sub:
        print("    该科 0 分按缺考不计入均分：%s" % "、".join(dropped_sub))
    print()
    print("  一、科目层（班级均分 vs 年级均分）")
    if subj_stats:
        print("     %-6s %-8s %-8s %-8s %s" % ("科目", "本班均分", "年级均分", "差距", "判定"))
        for s in sorted(subj_stats, key=lambda x: x["gap"]):
            tag = ("低于年级 %.0f 个百分点" % (-s["gap"] * 100)) if s["gap"] <= -GAP_MIN else (
                "高于年级 %.0f 个百分点" % (s["gap"] * 100) if s["gap"] >= GAP_MIN else "与年级持平")
            print("     %-6s %-8.1f %-8.1f %-+8.1f %s" % (
                s["h"], s["u_avg"], s["a_avg"], s["u_avg"] - s["a_avg"], tag))
    else:
        print("     成绩单里没认出科目列，科目层跳过。")

    # ---- 知识点层（要逐题得分）----
    kp_gaps, kp_good, kp_all = [], [], []
    sres = None
    if args.scores:
        bank = load_bank(args.bank)
        sres = load_scores(args.scores, bank)
        filled = attach_class(sres, grades)
        print()
        print_qc(sres["qc"], "逐题得分 %s，%s%s" % (
            sres["label"], sres["mode"], ("，按姓名补齐班级 %d 行" % filled) if filled else ""))
        rows = sres["rows"]
        u_rows = [r for r in rows if (r.get("cls") == unit if unit else True)]
        if not u_rows:
            print("     逐题得分里没有该班的数据（班级列缺失或名字对不上），知识点层跳过。")
        else:
            ku = {t["kp"]: t for t in kp_table(u_rows)}
            ka = {t["kp"]: t for t in kp_table(rows)}
            for kp, t in ku.items():
                a = ka.get(kp)
                g = t["P"] - a["P"] if a else 0.0
                row = {"kp": kp, "P": t["P"], "P_all": a["P"] if a else None,
                       "gap": g, "n_items": t["n_items"], "n_stu": t["n_students"],
                       "thin": t["n_items"] < KP_MIN_ITEMS}
                kp_all.append(row)
                if a and g <= -GAP_MIN:
                    kp_gaps.append(row)
                elif a and g >= GAP_MIN:
                    kp_good.append(row)
            if args.kp:                      # --kp：只盯一个知识点往下看
                keep = [k for k in kp_all if k["kp"] == args.kp]
                if not keep:
                    print()
                    print("  --kp %s：这次考试没考到这个点（逐题表里没有），知识点层不展开。"
                          % args.kp)
                kp_all[:] = keep
                kp_gaps[:] = [k for k in kp_gaps if k["kp"] == args.kp]
                kp_good[:] = [k for k in kp_good if k["kp"] == args.kp]
            print()
            print("  二、知识点层（本班得分率 vs 年级同知识点均值，差 ≥ %.0f 个百分点才说）%s" % (
                GAP_MIN * 100, ("｜只看 " + args.kp) if args.kp else ""))
            print("     %-14s %-6s %-8s %-8s %-8s %s" % (
                "知识点", "题数", "本班", "年级均", "差", "判定"))
            for k in sorted(kp_all, key=lambda x: x["gap"]):
                judge = ("没讲透" if k["gap"] <= -GAP_MIN else
                         ("讲法有效" if k["gap"] >= GAP_MIN else "持平"))
                if k["thin"]:
                    judge += "（仅 %d 道题，参考）" % k["n_items"]
                print("     %-14s %-6d %-8s %-8s %-+8.0f %s" % (
                    k["kp"], k["n_items"], pct(k["P"], 0),
                    "—" if k["P_all"] is None else pct(k["P_all"], 0),
                    k["gap"] * 100, judge))
            if not kp_gaps:
                print("     （没有知识点低于年级均值 %.0f 个百分点——这一块不用改讲法）" % (GAP_MIN * 100))
    else:
        print()
        print("  二、知识点层：没有 --scores 逐题得分，知识点粒度算不出来（只有科目总分）。")
        print("     要这一层，请给 --scores 逐题.csv（表头：学生,题号,知识点,满分,得分）。")

    # ---- 下次改进清单（三条，每条挂一个能核对的数字；不够就明说不够）----
    todo, used = [], set()

    def add(key, what, how, why):
        if key in used or len(todo) >= 3:
            return
        used.add(key)
        todo.append({"what": what, "how": how, "why": why})

    solid_gaps = [k for k in kp_gaps if not k["thin"]]
    thin_gaps = [k for k in kp_gaps if k["thin"]]
    for k in sorted(solid_gaps, key=lambda x: x["gap"])[:2]:
        add("kp:" + k["kp"],
            "%s：本班 %s vs 年级同点 %s（差 %+.0f 个百分点，%d 道题）" % (
                k["kp"], pct(k["P"], 0), pct(k["P_all"], 0), k["gap"] * 100, k["n_items"]),
            "下一节先花 10 分钟只讲这一个点：把定义重新落地（用数轴/实物，不复述教案），"
            "让学生用自己的话讲一遍，再上 2 道同点变式；先不碰综合题。",
            "同一批题对比出来的差距，说明这一块的具体讲法要重来，不是多做题能补的")
    for s in sorted([x for x in subj_stats if x["gap"] <= -GAP_MIN],
                    key=lambda x: x["gap"])[:2]:
        add("sub:" + s["h"],
            "%s：本班均分 %.1f，年级均分 %.1f（低 %.1f 分，%.0f 个百分点）" % (
                s["h"], s["u_avg"], s["a_avg"], s["a_avg"] - s["u_avg"], -s["gap"] * 100),
            "把这一科失分最集中的 2 个知识点各配一张「先问后答」的讲评卡，课上按卡问、"
            "不直接给答案；课后只留同点 3 道。",
            "分数差落在科目上，要改的是这一科的讲评次序与落点")
    for k in sorted(thin_gaps, key=lambda x: x["gap"])[:1]:
        add("kp:" + k["kp"],
            "%s：本班 %s vs 年级同点 %s（差 %+.0f 个百分点，只有 %d 道题）" % (
                k["kp"], pct(k["P"], 0), pct(k["P_all"], 0), k["gap"] * 100, k["n_items"]),
            "先在这个点上补 1～2 道变式题，把它变成「可对比的 2 道题」再下判断；"
            "眼下只把它放进下一节课的例题顺序里，不按整块重讲。",
            "题量太薄（%d 道），差距可能是这一道题本身的难度或歧义，先标出来不硬改" % k["n_items"])
    good = [k for k in kp_good if not k["thin"]] or kp_good
    if good:
        g = sorted(good, key=lambda x: -x["gap"])[0]
        add("good:" + g["kp"],
            "%s：本班 %s vs 年级同点 %s（高 %+.0f 个百分点，%d 道题）" % (
                g["kp"], pct(g["P"], 0), pct(g["P_all"], 0), g["gap"] * 100, g["n_items"]),
            "这一块的讲法已经跑通，固化成模板：把这次的引入问句与例题顺序写成一页「讲法卡」，"
            "下次同点直接复用；其余点照这个顺序改。",
            "有效的讲法要复用，别每次重新想")
    if subj_stats:
        low = min(subj_stats, key=lambda x: x["rate_u"])
        add("self:" + low["h"],
            "%s：本班得分率 %s（本班五科里最低；年级同科 %s，差 %+.0f 个百分点）" % (
                low["h"], pct(low["rate_u"], 0), pct(low["rate_a"], 0), low["gap"] * 100),
            "课时往这一科倾斜：讲评时先从这一科失分的那个步骤切入"
            "（要有 --scores 才能定位到具体题号），把这一步讲透再回综合题。",
            "这是本班五科之间的横向自比，不是与年级比——它说明课时该往哪儿倾斜")
    if len(todo) < 3:
        add("data", "数据缺口：%s" % (
                "没有逐题得分，知识点层只能看科目" if not args.scores
                else "低于均值(≥%s)的知识点只有 %d 个" % (pct(GAP_MIN), len(kp_gaps))),
            "补齐逐题得分（学生,题号,知识点,满分,得分）再跑一次；缺口没补前，"
            "不拿科目结论去推知识点讲法。",
            "样本不够就明说不足，不硬凑第三条")

    print()
    print("  三、下次改进清单（每条都挂一个数字依据；只谈知识点与讲法）")
    print("     所有文案在写出前过禁语自检（%d 个词），命中直接报错。" % len(FORBID_WORDS))
    body = []
    for i, t in enumerate(todo, 1):
        print("     %d) 依据：%s" % (i, t["what"]))
        print("        做法：%s" % t["how"])
        print("        为什么：%s" % t["why"])
        body.append(["%d" % i, t["what"], t["how"], t["why"]])
    checked = guard_text(*[x for t in todo for x in t.values()],
                         *[s for s in ("教学复盘", "本班", "年级")])
    print("     禁语自检：通过（查了 %d 个词，0 命中）" % checked)

    note = ("只看知识点与讲法，不评价老师个人。差距阈值 %.0f 个百分点；分析单位 %s（%d 人）。"
            % (GAP_MIN * 100, unit or "全表", len(unit_rows)))
    path = out_path("教学复盘_%s" % (unit or "全表"))
    sheets = [("改进清单", ["序号", "数字依据", "讲法建议", "为什么这么改"], body, note)]
    if kp_all:
        sheets.append(("知识点对比", ["知识点", "题数", "本班得分率", "年级同点得分率", "差(百分点)", "判定"],
                       [[k["kp"], k["n_items"], pct(k["P"], 1),
                         "—" if k["P_all"] is None else pct(k["P_all"], 1),
                         round(k["gap"] * 100, 1),
                         "没讲透" if k["gap"] <= -GAP_MIN else (
                             "讲法有效" if k["gap"] >= GAP_MIN else "持平")]
                        for k in sorted(kp_all, key=lambda x: x["gap"])], note))
    if subj_stats:
        sheets.append(("科目对比", ["科目", "本班均分", "年级均分", "差", "差(百分点)", "判定"],
                       [[s["h"], round(s["u_avg"], 1), round(s["a_avg"], 1),
                         round(s["u_avg"] - s["a_avg"], 1), round(s["gap"] * 100, 1),
                         "低于年级" if s["gap"] <= -GAP_MIN else (
                             "高于年级" if s["gap"] >= GAP_MIN else "持平")]
                        for s in sorted(subj_stats, key=lambda x: x["gap"])], None))
    sheets.append(("数据质检", ["项", "值"],
                   [["成绩单", "%s：%d 行 → 有效 %d 人" % (
                       grades["label"], grades["qc"]["rows_in"], grades["qc"]["rows_ok"])],
                    ["已排除", "、".join("%s(%s)" % (x["name"], x["why"])
                                        for x in grades["qc"]["excluded"]) or "无"],
                    ["逐题得分", "%s：%d 行 / %d 人 / %d 题（%s）" % (
                        sres["label"], sres["qc"]["rows_ok"], sres["qc"]["students"],
                        sres["qc"]["items"], sres["mode"]) if sres else "未提供"],
                    ["分析单位", "%s（%d 人）｜%s" % (unit or "全表", len(unit_rows), unit_note)],
                    ["禁语自检", "通过：%d 个评价老师个人的词，0 命中" % checked]],
                   None))
    write_xlsx(path, sheets)
    print()
    print("  → 已落表：%s" % path)
    return 0


# ══════════════════════════════════════════════════════════════════════════
def build_parser():
    p = argparse.ArgumentParser(
        description="教师工具 · 难点预判 / 连续下滑预警 / 教学复盘（0 次模型调用，数字全由代码算）")
    sub = p.add_subparsers(dest="cmd")

    a = sub.add_parser("hardkp", help="难点预判：这个知识点的第一道坎在哪")
    a.add_argument("--kp", help="知识点，例如 绝对值")
    a.add_argument("--chapter", help="章节号，例如 1.2")
    a.add_argument("--scores", help="逐题得分表（csv/xlsx），不给就只做题库侧判断")
    a.add_argument("--bank", default=None, help="题库.json（默认用教师工具/题库.json）")
    a.set_defaults(func=cmd_hardkp)

    b = sub.add_parser("drop", help="连续下滑预警：谁在稳步下滑、谁在大起大落")
    b.add_argument("--files", nargs="+", required=True, help="按时间顺序的成绩单 xlsx")
    b.set_defaults(func=cmd_drop)

    c = sub.add_parser("reflect", help="教学复盘：只谈知识点与讲法")
    c.add_argument("--files", nargs="+", required=True, help="成绩单 xlsx（给多份只看最近一次）")
    c.add_argument("--scores", help="逐题得分表，有它才做知识点层")
    c.add_argument("--class", dest="cls", default=None, help="班级，例如 九年级2班")
    c.add_argument("--kp", default=None, help="只聚焦某个知识点（可选）")
    c.add_argument("--bank", default=None, help="题库.json")
    c.set_defaults(func=cmd_reflect)
    return p


def main(argv=None):
    parser = build_parser()
    args = parser.parse_args(argv)
    if not getattr(args, "cmd", None):
        parser.print_help()
        return 2
    try:
        return args.func(args)
    except InsightError as e:
        print("！%s" % e)
        return 2


if __name__ == "__main__":
    sys.exit(main())
