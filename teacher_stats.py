
# -*- coding: utf-8 -*-
"""帮老师算成绩的小工具 —— 拖一张 Excel 进来，三十秒出全套统计。

用法：
    python teacher_stats.py 成绩表.xlsx
    python teacher_stats.py 成绩表.xlsx --full 120 --pass 72 --good 102

自动识别两种表：
    「每题」表：列名是 1,2,3...        -> 算每题得分率、总分的班级统计
    「每次」表：列名是 月考/期中...    -> 每次考试分别统计，不把分数相加

输出：同目录 <原名>_统计.xlsx
"""
import io, os, sys, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
WARN = PatternFill("solid", fgColor="FDE2E2")
GOOD = PatternFill("solid", fgColor="E4F2E4")
TITLE = Font(bold=True, size=13, color="1F4E79")

ID_KEYS = ("学号", "考号", "编号", "座号")
NAME_KEYS = ("姓名", "名字", "学生")
MISS = ("缺考", "缺", "免考", "-", "—", "/", "")

def to_num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip()
    if s in MISS: return None
    try: return float(re.sub(r"[^0-9.\-]", "", s))
    except Exception: return None

def find_header(ws):
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = [str(c.value).strip() if c.value is not None else "" for c in ws[r]]
        if any(any(k in v for k in NAME_KEYS + ID_KEYS) for v in vals):
            return r, vals
    return None, None

def head_style(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

def stat_of(vals, F, PL, GL):
    n = len(vals)
    if not n: return None
    avg = sum(vals) / n
    return {"n": n, "avg": avg, "max": max(vals), "min": min(vals),
            "pass": sum(1 for v in vals if v >= PL), "pass_rate": 100.0 * sum(1 for v in vals if v >= PL) / n,
            "good": sum(1 for v in vals if v >= GL), "good_rate": 100.0 * sum(1 for v in vals if v >= GL) / n}

def individual_mode(path, rows, cols, head, mode, F, n_valid, min_group):
    """有效人数 < min_group：不出任何群体结论，只出个体数据（原始分 + 各科得分率）。

    一对一 / 小班时，均分、及格率、优秀率、分数段这些指标描述的是「这个组」，样本都不够，
    算出来就是误导：1 个人也有 100% 及格率，3 个人也能算出一段「分布」。

    已验证：真实成绩单是「每次模式」（列名是 语文分数/数学分数…），本分支在真实数据上实测过
    2 人、1 人两种情况（产物含「个体数据 / 个体得分率」两个 sheet）。
    **未验证：每题模式（列名是 1,2,3…）下的小样本分支没有真实每题表可测**，只做过代码路径检查；
    没有为了点亮它去造数据。要覆盖这条路径，得先有一张真实的逐题得分表。
    """
    out = os.path.splitext(path)[0] + "_统计.xlsx"
    declare = "样本 %d 人，以下为个体数据，不作群体结论" % n_valid
    print()
    print("== 小样本模式 ==")
    print("   %s" % declare)
    print("   有效人数 %d < 群组下限 %d：不出群体统计（均分、及格率、优秀率、分数段）。"
          % (n_valid, min_group))
    wb = Workbook()
    ws = wb.active
    ws.title = "个体数据"
    ws["A1"] = declare
    ws["A1"].font = Font(bold=True, size=13, color="C00000")
    ws["A2"] = ("有效人数 %d < 群组下限 %d：一个组凑不出群体口径，所以这里只列原始分数。"
                "需要群体统计请补足人数，或用 --min-group 调整下限。" % (n_valid, min_group))
    hdr = ["学号", "姓名"] + [head[i] for i in cols] + (["总分"] if mode == "每题" else ["平均分"])
    ws.append([])
    ws.append(hdr)
    for c in range(1, len(hdr) + 1):
        cell = ws.cell(row=3, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center")
    for st in rows:
        with_scores = [st["sc"].get(head[i], "") for i in cols]
        if mode == "每题":
            tail = sum(st["sc"].values())
        else:
            vs = list(st["sc"].values())
            tail = round(sum(vs) / len(vs), 1) if vs else ""
        ws.append([st["id"], st["name"]] + with_scores + [tail])
    for i, w in enumerate([10, 10] + [8] * len(cols) + [10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("个体得分率")
    ws2["A1"] = declare
    ws2["A1"].font = Font(bold=True, size=13, color="C00000")
    ws2["A2"] = "得分率 = 该栏得分 ÷ 满分（%.0f）" % F
    hdr2 = ["学号", "姓名"] + ["%s得分率" % head[i] for i in cols] + ["平均得分率"]
    ws2.append([])
    ws2.append(hdr2)
    for c in range(1, len(hdr2) + 1):
        cell = ws2.cell(row=3, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center")
    for st in rows:
        rates = []
        for i in cols:
            v = st["sc"].get(head[i])
            rates.append("" if v is None else "%.0f%%" % (100.0 * v / F))
        got = [st["sc"][head[i]] for i in cols if head[i] in st["sc"]]
        avg_rate = "%.0f%%" % (100.0 * sum(got) / (F * len(got))) if got else ""
        ws2.append([st["id"], st["name"]] + rates + [avg_rate])
    for i, w in enumerate([10, 10] + [12] * len(cols) + [12], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    wb.save(out)
    print("   个体数据已写入: %s（sheet：个体数据 / 个体得分率）" % out)
    return out


def analyze(path, full=None, passline=None, goodline=None, min_group=3):
    ws = load_workbook(path, data_only=True).active
    hrow, head = find_header(ws)
    if not hrow: raise SystemExit("没找到表头：需要有「学号」或「姓名」这一列")

    id_col = next((i for i, h in enumerate(head) if any(k in h for k in ID_KEYS)), None)
    name_col = next((i for i, h in enumerate(head) if any(k in h for k in NAME_KEYS)), None)
    cols = [i for i, h in enumerate(head) if i not in (id_col, name_col) and h]

    def serial(h): return bool(re.fullmatch(r"\d{1,3}", str(h).strip()))
    mode = "每题" if sum(1 for i in cols if serial(head[i])) >= max(3, len(cols) * 0.6) else "每次"

    rows = []
    for r in range(hrow + 1, ws.max_row + 1):
        vals = [c.value for c in ws[r]]
        if all(v is None or str(v).strip() == "" for v in vals): continue
        sid = vals[id_col] if id_col is not None and id_col < len(vals) else ""
        nm = vals[name_col] if name_col is not None and name_col < len(vals) else ""
        if not nm and not sid: continue
        sc = {}
        for i in cols:
            if i < len(vals):
                v = to_num(vals[i])
                if v is not None: sc[head[i]] = v
        rows.append({"id": sid, "name": nm, "sc": sc})
    if not rows: raise SystemExit("表里没有学生数据")

    allv = [v for st in rows for v in st["sc"].values()]
    F = full or (max(allv) if allv else 100)
    PL = passline if passline is not None else F * 0.6
    GL = goodline if goodline is not None else F * 0.85

    print("表头: 第%d行 ｜ 学号列:%s ｜ 姓名列:%s" % (hrow, head[id_col] if id_col is not None else "无",
                                                   head[name_col] if name_col is not None else "无"))
    print("数据列 %d 个 ｜ 判定【%s】｜ 满分 %.0f ｜ 及格 %.0f ｜ 优秀 %.0f" % (len(cols), mode, F, PL, GL))
    print("学生 %d 人" % len(rows))

    n_valid = sum(1 for st in rows if st["sc"] and sum(st["sc"].values()) > 0)
    if n_valid < min_group:
        return individual_mode(path, rows, cols, head, mode, F, n_valid, min_group)

    out = os.path.splitext(path)[0] + "_统计.xlsx"
    wb = Workbook()

    if mode == "每题":
        for st in rows:
            st["total"] = sum(st["sc"].values()); st["n"] = len(st["sc"])
        totals = [st["total"] for st in rows]
        ranked = sorted(rows, key=lambda x: -x["total"])
        for i, st in enumerate(ranked, 1): st["rank"] = i
        cls = stat_of(totals, F, PL, GL)

        ws1 = wb.active; ws1.title = "班级概览"
        ws1["A1"] = "班级概览（每题模式：%d 道题，满分 %d）" % (len(cols), F); ws1["A1"].font = TITLE
        items = [("学生人数", len(rows)), ("平均分", round(cls["avg"], 1)), ("最高分", cls["max"]),
                 ("最低分", cls["min"]), ("及格率", "%.0f%%（%d 人）" % (cls["pass_rate"], cls["pass"])),
                 ("优秀率", "%.0f%%（%d 人）" % (cls["good_rate"], cls["good"]))]
        rr = 3
        for k, v in items:
            ws1.cell(row=rr, column=1, value=k).font = Font(bold=True)
            ws1.cell(row=rr, column=2, value=v); rr += 1
        ws1.column_dimensions["A"].width = 14; ws1.column_dimensions["B"].width = 22

        ws2 = wb.create_sheet("每人明细")
        ws2.append(["名次", "学号", "姓名", "总分"] + [head[i] for i in cols])
        head_style(ws2, 4 + len(cols))
        for st in ranked:
            ws2.append([st["rank"], st["id"], st["name"], st["total"]] + [st["sc"].get(head[i], "") for i in cols])
            mark = WARN if st["total"] < PL else (GOOD if st["total"] >= GL else None)
            if mark: ws2.cell(row=ws2.max_row, column=4).fill = mark
        for i, w in enumerate([6, 10, 10, 8] + [7] * len(cols), 1):
            ws2.column_dimensions[get_column_letter(i)].width = w

        ws3 = wb.create_sheet("每题正确率")
        ws3.append(["题号", "作答人数", "平均得分", "得分率", "判定"])
        head_style(ws3, 5)
        for i in cols:
            h = head[i]
            vs = [st["sc"][h] for st in rows if h in st["sc"]]
            if not vs: continue
            rate = sum(vs) / (F * len(vs)) * 100
            tag = "白给分" if rate >= 90 else ("太难" if rate <= 20 else "")
            ws3.append([h, len(vs), round(sum(vs) / len(vs), 2), "%d%%" % round(rate), tag])
            if tag: ws3.cell(row=ws3.max_row, column=4).fill = WARN
        for i, w in enumerate([8, 10, 10, 9, 10], 1):
            ws3.column_dimensions[get_column_letter(i)].width = w

        segs = []
        step = max(1, round(F / 6 / 5) * 5)
        for k in range(6):
            a = max(0, F - step * (k + 1)); b = F - step * k
            segs.append((a, b, sum(1 for t in totals if a < t <= b)))
        segs.reverse()
        ws4 = wb.create_sheet("分数段分布")
        ws4.append(["分数段", "人数", "占比", "图示"]); head_style(ws4, 4)
        for a, b, c in segs:
            ws4.append(["%d–%d" % (a, b), c, "%.0f%%" % (100 * c / len(totals)), "█" * c])
        for i, w in enumerate([12, 8, 8, 40], 1):
            ws4.column_dimensions[get_column_letter(i)].width = w

        print()
        print("== 班级概览 ==")
        for k, v in items: print("   %-6s %s" % (k, v))
        print("== 分数段 ==")
        for a, b, c in segs: print("   %3d–%3d  %2d 人  %s" % (a, b, c, "█" * c))

    else:
        # 每次考试分别统计，绝不把分数相加
        per_exam = []
        for i in cols:
            h = head[i]
            vs = [st["sc"][h] for st in rows if h in st["sc"]]
            s = stat_of(vs, F, PL, GL)
            if s: s["name"] = h; per_exam.append(s)
        for st in rows:
            vs = list(st["sc"].values())
            st["avg"] = round(sum(vs) / len(vs), 1) if vs else 0
            st["n"] = len(vs)
        ranked = sorted(rows, key=lambda x: -x["avg"])
        for i, st in enumerate(ranked, 1): st["rank"] = i
        clsa = stat_of([st["avg"] for st in rows], F, PL, GL)

        ws1 = wb.active; ws1.title = "各次考试对比"
        ws1.append(["考试", "人数", "平均分", "最高", "最低", "及格率", "优秀率"])
        head_style(ws1, 7)
        for s in per_exam:
            ws1.append([s["name"], s["n"], round(s["avg"], 1), s["max"], s["min"],
                        "%.0f%%" % s["pass_rate"], "%.0f%%" % s["good_rate"]])
        ws1.append(["【平均】", len(rows), clsa["avg"], clsa["max"], clsa["min"],
                    "%.0f%%" % clsa["pass_rate"], "%.0f%%" % clsa["good_rate"]])
        for c in range(1, 8): ws1.cell(row=ws1.max_row, column=c).font = Font(bold=True)
        for i, w in enumerate([16, 8, 10, 8, 8, 10, 10], 1):
            ws1.column_dimensions[get_column_letter(i)].width = w

        ws2 = wb.create_sheet("每人明细")
        ws2.append(["名次", "学号", "姓名"] + [head[i] for i in cols] + ["平均分", "缺考次数"])
        head_style(ws2, 5 + len(cols))
        for st in ranked:
            miss = len(cols) - st["n"]
            ws2.append([st["rank"], st["id"], st["name"]] + [st["sc"].get(head[i], "缺考") for i in cols] +
                       [st["avg"], miss])
            mark = WARN if st["avg"] < PL else (GOOD if st["avg"] >= GL else None)
            if mark: ws2.cell(row=ws2.max_row, column=3 + len(cols) + 1).fill = mark
        for i, w in enumerate([6, 10, 10] + [12] * len(cols) + [9, 9], 1):
            ws2.column_dimensions[get_column_letter(i)].width = w
        ws2.auto_filter.ref = "A1:%s%d" % (get_column_letter(5 + len(cols)), ws2.max_row)

        avgs = [st["avg"] for st in rows]
        step = max(1, round(F / 6 / 5) * 5)
        segs = []
        for k in range(6):
            a = max(0, F - step * (k + 1)); b = F - step * k
            segs.append((a, b, sum(1 for t in avgs if a < t <= b)))
        segs.reverse()
        ws3 = wb.create_sheet("分数段分布")
        ws3.append(["分数段（按个人平均分）", "人数", "占比", "图示"]); head_style(ws3, 4)
        for a, b, c in segs:
            ws3.append(["%d–%d" % (a, b), c, "%.0f%%" % (100 * c / len(avgs)), "█" * c])
        for i, w in enumerate([22, 8, 8, 40], 1):
            ws3.column_dimensions[get_column_letter(i)].width = w

        print()
        print("== 各次考试 ==")
        print("   %-14s %5s %7s %7s %8s %8s" % ("考试", "人数", "平均", "最高", "及格率", "优秀率"))
        for s in per_exam:
            print("   %-14s %5d %7.1f %7.0f %7.0f%% %7.0f%%" % (s["name"], s["n"], s["avg"], s["max"], s["pass_rate"], s["good_rate"]))
        print("   %-14s %5d %7.1f %7.0f %7.0f%% %7.0f%%" % ("【个人平均】", len(rows), clsa["avg"], clsa["max"], clsa["pass_rate"], clsa["good_rate"]))
        print()
        print("== 个人平均分分布 ==")
        for a, b, c in segs: print("   %3d–%3d  %2d 人  %s" % (a, b, c, "█" * c))

    wb.save(out)
    print()
    print("已生成:", out)
    return out

if __name__ == "__main__":
    kw = {}
    for i, a in enumerate(sys.argv):
        if a == "--full": kw["full"] = float(sys.argv[i + 1])
        if a == "--pass": kw["passline"] = float(sys.argv[i + 1])
        if a == "--good": kw["goodline"] = float(sys.argv[i + 1])
        if a == "--min-group": kw["min_group"] = int(sys.argv[i + 1])
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        print("用法: python teacher_stats.py 成绩表.xlsx [--full 120] [--pass 72] [--good 102] [--min-group 3]")
        sys.exit(0)
    analyze(args[0], **kw)
