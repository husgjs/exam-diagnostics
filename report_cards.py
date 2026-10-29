# -*- coding: utf-8 -*-
"""教师工具 · 一页纸三件套（纯标准库 + openpyxl，0 次模型调用）

    python report_cards.py student --name 张三        # 一人一页学情卡
    python report_cards.py class   --name 九年级2班    # 一页纸班级诊断
    python report_cards.py talk    --name 张三        # 教师话术卡

口径（全部由代码算，模型一次都不调）：
  · 满分：语文 100 / 数学 150 / 英语 150 / 物理 100 / 化学 100，总分 600
  · 强弱科按「得分率 = 分数 ÷ 满分」比，不按原始分（五科满分不同，
    按原始分比，数学永远最强、语文永远最弱，是假的）
  · 及格 = 360 分（600 的 60%）；优秀 = 510 分（600 的 85%）
  · 「两极分化」判定：班级标准差 ÷ 年级标准差 ≥ 1.15（≤ 0.85 判比年级更集中）
  · 所有建议里的「如果再……会怎样」都是拿全年级的真实分布重排一遍算出来的，
    不是估的
"""
import argparse
import html
import os
import statistics
import sys
import unicodedata

import openpyxl

try:                      # Windows 控制台/管道默认 GBK，中文会变成乱码，这里统一成 UTF-8
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
XLSX = os.path.join(ROOT, "九年级月考成绩单.xlsx")
OUT_DIR = os.path.join(HERE, "输出")

SUBJECTS = (("语文", "语文分数", 100), ("数学", "数学分数", 150), ("英语", "英语分数", 150),
            ("物理", "物理分数", 100), ("化学", "化学分数", 100))
SUBJ_NAMES = [s[0] for s in SUBJECTS]
FULL = {name: full for name, _col, full in SUBJECTS}
TOTAL_FULL = 600
PASS_LINE = 360        # 600 × 60%
GOOD_LINE = 510        # 600 × 85%
# 「两极分化」不设绝对分数线：拿班级标准差去比年级标准差。
# 一开始用的绝对阈值（σ ≥ 60 分）在真实数据上八个班全部触发，等于没判；
# 相对判据才能区分「这个班比年级更散」和「这个班比年级更齐」。
MIN_GROUP = 3          # 群组下限：有效人数低于它就不出群体结论（一对一/小班）
SPLIT_RATIO = 1.15     # σ_班 / σ_年级 ≥ 1.15 → 两极分化
TIGHT_RATIO = 0.85     # σ_班 / σ_年级 ≤ 0.85 → 比年级更集中
NEAR_BAND = 30         # 临界生：距及格线 ±30 分
NEAR_MAX = 8           # 一页最多盯 8 个人
BIAS_PP = 15.0         # 科目极差超过 15 个百分点算偏科

# 页面里我们自己写的分析文字，一律不许出现这些词（空话 / 预测中考）。
# 「不要说的话」那一节是反面例句，单独放行，见 _guard()。
BANNED = ("继续努力", "再接再厉", "加油", "你很聪明", "很聪明", "有潜力", "潜力大",
          "态度不端正", "态度有问题", "不认真", "不够努力", "中考", "考上", "稳上",
          "提前录取", "一定能", "只要努力", "静待花开")


class UserError(Exception):
    """给老师看的提示（查无此人 / 查无此班 / 文件没找到）。只打印一句话，不打异常栈。"""


# ------------------------------------------------------------------ 读数据
def load_rows(path=None):
    path = path or XLSX
    if not os.path.exists(path):
        raise UserError("找不到成绩单：%s" % path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise UserError("成绩单打不开：%s（%s）" % (os.path.basename(path), str(e)[:80]))
    ws = wb["九年级"] if "九年级" in wb.sheetnames else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        head = [str(c or "").strip() for c in next(rows_iter)]
    except StopIteration:
        raise UserError("成绩单是空的：%s" % os.path.basename(path))
    idx = {}
    for i, h in enumerate(head):
        if h and h not in idx:
            idx[h] = i
    need = ["班级", "姓名", "总分"] + [c for _n, c, _f in SUBJECTS]
    miss = [x for x in need if x not in idx]
    if miss:
        raise UserError("成绩单缺少这些列：%s（表头是：%s）" % ("、".join(miss), "、".join(h for h in head if h)))

    def num(row, col, default=0.0):
        try:
            v = row[idx[col]]
        except (IndexError, KeyError):
            return default
        if isinstance(v, bool) or v is None:
            return default
        try:
            return float(v)
        except (TypeError, ValueError):
            return default

    out = []
    for row in rows_iter:
        if not row:
            continue
        try:
            name = str(row[idx["姓名"]] or "").strip()
        except IndexError:
            continue
        if not name:
            continue
        scores = {nm: num(row, col) for nm, col, _f in SUBJECTS}
        try:
            cls = str(row[idx["班级"]] or "").strip()
        except IndexError:
            cls = ""
        total_col = num(row, "总分", None)
        total = total_col if total_col is not None else sum(scores.values())
        mini = scores["语文"] + scores["数学"] + scores["英语"]
        # 总分 0 = 缺考/未录入：人还在班里（占班级人数、想进临界生名单也能进），
        # 但不进均分、不进标准差、不排名次 —— 6 个 0 分会把年级均分拉低 7 分。
        out.append({"name": name, "cls": cls, "scores": scores, "total": total, "mini": mini,
                    "absent": total <= 1e-9})
    if not out:
        raise UserError("成绩单里一条学生记录都没读到：%s" % os.path.basename(path))
    return out


def valid_rows(rows):
    """计入统计与名次的行：总分 0（缺考/未录入）不算。"""
    return [r for r in rows if not r["absent"]]


def grade_index(rows):
    """年级口径：均分、标准差、各科均分一律只算有效成绩（剔除缺考）。"""
    ok = valid_rows(rows)
    per = {}
    for nm in SUBJ_NAMES:
        vals = [r["scores"][nm] for r in ok]
        per[nm] = {"mean": statistics.mean(vals), "sigma": statistics.pstdev(vals)}
    totals = [r["total"] for r in ok]
    return {"subjects": per, "n": len(ok), "n_all": len(rows), "n_absent": len(rows) - len(ok),
            "mean": statistics.mean(totals), "sigma": statistics.pstdev(totals),
            "pass": sum(1 for t in totals if t >= PASS_LINE),
            "good": sum(1 for t in totals if t >= GOOD_LINE),
            "max": max(totals)}


def _ranks(values):
    """竞赛式排名：同分同名次（1,1,3）。"""
    order = sorted(range(len(values)), key=lambda i: -values[i])
    ranks = [0] * len(values)
    prev_v, prev_rank = None, 0
    for pos, i in enumerate(order, 1):
        v = values[i]
        if prev_v is not None and abs(v - prev_v) <= 1e-9:
            ranks[i] = prev_rank
        else:
            ranks[i] = pos
            prev_rank = pos
        prev_v = v
    return ranks


def with_ranks(rows):
    """补 年级名次 / 班级名次 / 语数外小总分名次。
    名次只在有效成绩之间排：缺考的人名次留 None，不占别人的名次。"""
    ok = valid_rows(rows)
    for r in rows:
        r["grade_rank"] = r["mini_rank"] = r["class_rank"] = None
    for r, rk in zip(ok, _ranks([r["total"] for r in ok])):
        r["grade_rank"] = rk
    for r, rk in zip(ok, _ranks([r["mini"] for r in ok])):
        r["mini_rank"] = rk
    by_cls = {}
    for r in rows:
        by_cls.setdefault(r["cls"], []).append(r)
    for members in by_cls.values():
        v = valid_rows(members)
        for r, rk in zip(v, _ranks([r["total"] for r in v])):
            r["class_rank"] = rk
        for r in members:
            r["class_size"] = len(members)    # 班里实际几个人（含缺考）
            r["class_valid"] = len(v)         # 计入名次和统计的几个人
    return rows


def rank_if(rows, total):
    """这个总分在年级里排第几（只跟有效成绩比，同分同名次）。"""
    return 1 + sum(1 for r in valid_rows(rows) if r["total"] > total + 1e-9)


# ------------------------------------------------------------------ 小工具
def _w(s):
    """显示宽度：中文算 2 列，用来在纯文本版里对齐表格。"""
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in str(s))


def _pad(s, width, right=False):
    s = str(s)
    gap = " " * max(0, width - _w(s))
    return gap + s if right else s + gap


def rank_note(rows, before, total):
    """名次真的没动时说清楚，别写「1 → 1」让人以为算错了。"""
    if rank_if(rows, total) != before:
        return ""
    return "（名次没有变化：现在就是第 %d 名，这一项体现在分数上）" % before


def f1(x):
    return "%.1f" % x


def fg(x):
    return ("%d" % round(x)) if abs(x - round(x)) < 1e-9 else ("%.1f" % x)


def rate(score, full):
    return 100.0 * score / full if full else 0.0


def signed(x, nd=1):
    return ("+" if x >= 0 else "-") + ("%.*f" % (nd, abs(x)))


def _guard(*texts):
    """我们自己的分析文字里不许出现空话/预测；出现就是代码写错了，直接报错不发布。"""
    bad = []
    for t in texts:
        for w in BANNED:
            if w in str(t):
                bad.append(w)
    if bad:
        raise RuntimeError("页面里出现了不该出现的说法：%s" % "、".join(sorted(set(bad))))


def pick_student(rows, name, cls=None):
    name = (name or "").strip()
    hit = [r for r in rows if r["name"] == name and (not cls or r["cls"] == cls)]
    if not hit:
        sample = "、".join(r["name"] for r in rows[:3])
        if cls:
            raise UserError("没有找到这个学生：「%s」（班级：%s）。同名同班都没有记录，"
                            "请核对姓名是否与成绩单一致（成绩单里的名字例如：%s）。" % (name, cls, sample))
        raise UserError("没有找到这个学生：「%s」。请核对姓名是否与成绩单一致"
                        "（成绩单里的名字例如：%s）。" % (name, sample))
    if len(hit) > 1:
        raise UserError("有 %d 个学生都叫「%s」：%s。请加 --class 指定班级。"
                        % (len(hit), name, "、".join(sorted(r["cls"] for r in hit))))
    return hit[0]


def pick_class(rows, cls):
    cls = (cls or "").strip()
    names = sorted({r["cls"] for r in rows if r["cls"]})
    if not cls or cls not in names:
        raise UserError("没有找到这个班：「%s」。现有班级：%s。"
                        % (cls or "（空）", "、".join(names)))
    members = [r for r in rows if r["cls"] == cls]
    if not members:
        raise UserError("「%s」这个班没有学生记录。" % cls)
    return members


# ------------------------------------------------------------------ #66 一人一页学情卡
def build_student_card(rows, gidx, stu):
    if stu["absent"]:
        raise UserError("「%s」（%s）本次总分 0，按缺考/未录入处理：没有可用的成绩数据，"
                        "均差、名次、建议都算不出来。请先补录成绩再出学情卡。" % (stu["name"], stu["cls"]))
    cls = stu["cls"]
    rates = {nm: rate(stu["scores"][nm], FULL[nm]) for nm in SUBJ_NAMES}
    strong = max(SUBJ_NAMES, key=lambda nm: (rates[nm], FULL[nm]))
    weak = min(SUBJ_NAMES, key=lambda nm: (rates[nm], FULL[nm]))
    spread = rates[strong] - rates[weak]

    table = [{"head": ["学科", "分数", "满分", "得分率", "年级均分", "年级均差"],
              "rows": []}]
    for nm in SUBJ_NAMES:
        s = stu["scores"][nm]
        gm = gidx["subjects"][nm]["mean"]
        table[0]["rows"].append([nm, fg(s), fg(FULL[nm]), f1(rates[nm]) + "%", f1(gm), signed(s - gm)])
    table[0]["rows"].append(["总分", fg(stu["total"]), fg(TOTAL_FULL),
                             f1(rate(stu["total"], TOTAL_FULL)) + "%", f1(gidx["mean"]),
                             signed(stu["total"] - gidx["mean"])])

    advices = []
    # 建议 1：最弱科 —— 拿全年级分布把「补上去之后排第几」算出来
    w_rate, w_score, w_full = rates[weak], stu["scores"][weak], FULL[weak]
    g_rate = rate(gidx["subjects"][weak]["mean"], w_full)
    if w_rate < g_rate:
        target_score = gidx["subjects"][weak]["mean"]
        gain = target_score - w_score
        would = stu["total"] + gain
        advices.append("补最弱科%s：得分率 %s%%（年级 %s%%），差 %s 个百分点，"
                       "折成分数是 %s 分（%s → 年级均分 %s 分）。"
                       "把这 %s 分补回来，总分 %s → %s，年级名次 %d → %d%s。"
                       "动作：%s 每周 2 次、每次 20 分钟，只做同一种题型，错一道当场重做一道。"
                       % (weak, f1(w_rate), f1(g_rate), f1(g_rate - w_rate), f1(gain),
                          fg(w_score), f1(target_score), f1(gain), fg(stu["total"]), f1(would),
                          stu["grade_rank"], rank_if(rows, would),
                          rank_note(rows, stu["grade_rank"], would), weak))
    else:
        gain = w_full * 0.05
        would = stu["total"] + gain
        advices.append("最弱科%s已经高于年级均分：得分率 %s%%（年级 %s%%），高 %s 个百分点。"
                       "再提 5 个百分点是 %s 分，总分 %s → %s，年级名次 %d → %d%s。"
                       "动作：%s 每周 2 次、每次 20 分钟，只做同一种题型，错一道当场重做一道。"
                       % (weak, f1(w_rate), f1(g_rate), f1(w_rate - g_rate), f1(gain),
                          fg(stu["total"]), f1(would), stu["grade_rank"], rank_if(rows, would),
                          rank_note(rows, stu["grade_rank"], would), weak))

    # 建议 2：最强科 —— 算「这科掉回年级均分会掉多少名」，说明它的分量
    s_rate, s_score = rates[strong], stu["scores"][strong]
    gs_rate = rate(gidx["subjects"][strong]["mean"], FULL[strong])
    lead = s_score - gidx["subjects"][strong]["mean"]
    would2 = stu["total"] - lead
    lead_pp = s_rate - gs_rate
    advices.append("保住最强科%s：得分率 %s%%（年级 %s%%），领先 %s 个百分点，"
                   "折成分数是 %s 分。这科要是掉回年级均分，总分 %s → %s，年级名次 %d → %d —— %s"
                   "动作：%s 不再加新题量，每周 1 次限时成套卷；从%s挪出来的时间全部给%s。"
                   % (strong, f1(s_rate), f1(gs_rate), f1(lead_pp), f1(lead),
                      fg(stu["total"]), f1(would2), stu["grade_rank"], rank_if(rows, would2),
                      ("它是你现在名次的主要支撑。" if lead_pp >= 5 else
                       "领先不到 5 个百分点，保它的意义是别让名次再往下走，别指望它拉分。"),
                      strong, strong, weak))

    # 建议 3：结构 —— 偏科就补短板，均衡就同步抬
    if spread >= BIAS_PP:
        gain3 = (s_rate - w_rate) / 100.0 * w_full
        would3 = stu["total"] + gain3
        advices.append("科目不均衡：最强%s %s%% 与最弱%s %s%% 差 %s 个百分点"
                       "（判定阈值 %s 个百分点）。把%s拉到%s的水平，总分 %s → %s，年级名次 %d → %d%s。"
                       "动作：本月只补这一个缺口，别的科目维持现有练习量。"
                       % (strong, f1(s_rate), weak, f1(w_rate), f1(spread), fg(BIAS_PP),
                          weak, strong, fg(stu["total"]), f1(would3), stu["grade_rank"],
                          rank_if(rows, would3), rank_note(rows, stu["grade_rank"], would3)))
    else:
        gain3 = sum(FULL[nm] for nm in SUBJ_NAMES) * 0.03
        would3 = stu["total"] + gain3
        dmini = stu["mini_rank"] - stu["grade_rank"]
        advices.append("科目均衡：最强%s %s%% 与最弱%s %s%% 只差 %s 个百分点"
                       "（判定阈值 %s 个百分点），拉不开差距，也不用拆东墙补西墙。"
                       "五科各提 3 个百分点合计 %s 分，总分 %s → %s，年级名次 %d → %d%s。"
                       "旁证：语数外小总分名次第 %d、总分名次第 %d，%s。"
                       "动作：五科每周各 1 次 15 分钟小测，只看错题原因，不刷整套卷。"
                       % (strong, f1(s_rate), weak, f1(w_rate), f1(spread), fg(BIAS_PP),
                          f1(gain3), fg(stu["total"]), f1(would3), stu["grade_rank"],
                          rank_if(rows, would3), rank_note(rows, stu["grade_rank"], would3),
                          stu["mini_rank"], stu["grade_rank"],
                          ("物理化学把小总分名次又往前推了 %d 名" % dmini) if dmini > 0
                          else ("物理化学把小总分名次往后拉了 %d 名" % (-dmini) if dmini < 0
                                else "物理化学没有改变名次")))
    _guard(*advices)

    return {
        "kind": "student",
        "title": "学情卡 ｜ %s %s" % (cls, stu["name"]),
        "sub": "九年级月考 · 满分 600 分（语文 100 / 数学 150 / 英语 150 / 物理 100 / 化学 100）",
        "big": [("总分", fg(stu["total"])),
                ("年级名次", "%d / %d" % (stu["grade_rank"], gidx["n"])),
                # 本班有效人数不足群组下限时，班内名次没有意义：一对一报告不拿「1/1」当结论
                ("班级名次", ("%d / %d" % (stu["class_rank"], stu["class_valid"]))
                 if stu["class_valid"] >= MIN_GROUP
                 else "—（本班有效 %d 人，少于 %d 人，不排名次）" % (stu["class_valid"], MIN_GROUP)),
                ("总分得分率", f1(rate(stu["total"], TOTAL_FULL)) + "%")],
        "tables": table,
        "strong": "%s %s%%（年级 %s%%）" % (strong, f1(s_rate), f1(gs_rate)),
        "weak": "%s %s%%（年级 %s%%）" % (weak, f1(w_rate), f1(g_rate)),
        "spread": f1(spread) + " 个百分点",
        "advices": advices,
        "notes": ["强弱科按得分率（分数 ÷ 满分）比，不按原始分——五科满分不同，"
                  "按原始分比数学永远最强、语文永远最弱。",
                  "年级名次、年级均分、年级均差都按剔除缺考后的 %d 人算"
                  "（全年级 %d 人，其中 %d 人总分 0 = 缺考/未录入，不计入统计与名次）。"
                  % (gidx["n"], gidx["n_all"], gidx["n_absent"]),
                  "班级名次在本班计入排名的 %d 人里排（本班共 %d 人，%d 人缺考不占名次）。"
                  % (stu["class_valid"], stu["class_size"], stu["class_size"] - stu["class_valid"])
                  if stu["class_valid"] >= MIN_GROUP
                  else "本班有效成绩只有 %d 人（少于群组下限 %d 人）：不排名次，也不做横向比较。"
                       % (stu["class_valid"], MIN_GROUP)],
        "file": "学情卡_%s.html" % stu["name"],
    }


# ------------------------------------------------------------------ #13 一页纸班级诊断
def build_class_card(rows, gidx, cls, min_group=None):
    if min_group is None:
        min_group = MIN_GROUP
    members = pick_class(rows, cls)
    ok = valid_rows(members)                      # 计入统计的人（缺考剔除）
    totals = [r["total"] for r in ok]
    n = len(totals)
    if n < min_group:
        # 拒绝生成，不静默出空卡：1~2 个人撑不起班级诊断的任何一条群体结论
        raise UserError(
            "「%s」有效成绩只有 %d 人，少于群组下限 %d 人 —— 一页纸班级诊断要出的是群体结论"
            "（平均水平、达标比例、班内差异幅度判断），这个人数算出来就是误导，所以不生成。"
            "要看单人请用 student 子命令；确实要按这个小组出诊断，请显式把下限调小（--min-group %d）。"
            % (cls, n, min_group, n if n > 0 else 1))
    mean = statistics.mean(totals)
    sigma = statistics.pstdev(totals)
    pass_n = sum(1 for t in totals if t >= PASS_LINE)
    good_n = sum(1 for t in totals if t >= GOOD_LINE)
    pass_rate = 100.0 * pass_n / n
    good_rate = 100.0 * good_n / n
    zero = [r for r in members if r["absent"]]      # 与统计口径保持一致：总分 0 = 缺考

    low = mean < PASS_LINE
    ratio = (sigma / gidx["sigma"]) if gidx["sigma"] else 0.0
    if low:
        head = ("整体差：均分 %s 分，比及格线 %d 分低 %s 分" % (f1(mean), PASS_LINE, f1(PASS_LINE - mean)))
    else:
        head = ("不是整体差：均分 %s 分，比及格线 %d 分高 %s 分" % (f1(mean), PASS_LINE, f1(mean - PASS_LINE)))
    if ratio >= SPLIT_RATIO:
        tail = ("，而且是两极分化 —— 标准差 %s 分，是年级（%s 分）的 %.2f 倍，超过 %.2f 倍阈值，"
                "头部和尾部隔得很开。" % (f1(sigma), f1(gidx["sigma"]), ratio, SPLIT_RATIO))
    elif ratio <= TIGHT_RATIO:
        tail = ("，也不是两极分化 —— 标准差 %s 分，是年级（%s 分）的 %.2f 倍，不到 %.2f 倍，"
                "班内差距比年级小。" % (f1(sigma), f1(gidx["sigma"]), ratio, TIGHT_RATIO))
    else:
        tail = ("，也不是两极分化 —— 标准差 %s 分，是年级（%s 分）的 %.2f 倍，落在 %.2f–%.2f 倍之间，"
                "班内差距与年级相当。" % (f1(sigma), f1(gidx["sigma"]), ratio, TIGHT_RATIO, SPLIT_RATIO))
    verdict = head + tail

    threshold = ("判断阈值：班级标准差 ÷ 年级标准差（%s 分）≥ %.2f 判「两极分化」，"
                 "≤ %.2f 判「班内差距比年级小」，中间判「与年级相当」。"
                 "均分是否过 %d 分及格线（满分 %d 的 60%%）决定算不算「整体差」；"
                 "优秀线 %d 分（85%%）。"
                 % (f1(gidx["sigma"]), SPLIT_RATIO, TIGHT_RATIO, PASS_LINE, TOTAL_FULL, GOOD_LINE))

    subject_rows = []
    rates = {}
    for nm in SUBJ_NAMES:
        cm = statistics.mean([r["scores"][nm] for r in ok])
        gm = gidx["subjects"][nm]["mean"]
        rates[nm] = rate(cm, FULL[nm])
        subject_rows.append([nm, f1(cm), fg(FULL[nm]), f1(rates[nm]) + "%", f1(gm),
                             signed(cm - gm), f1(rate(gm, FULL[nm])) + "%"])
    strong = max(SUBJ_NAMES, key=lambda nm: (rates[nm], FULL[nm]))
    weak = min(SUBJ_NAMES, key=lambda nm: (rates[nm], FULL[nm]))

    near = sorted([r for r in members if abs(r["total"] - PASS_LINE) <= NEAR_BAND],
                  key=lambda r: (abs(r["total"] - PASS_LINE), -r["total"]))
    watch = []
    for r in near[:NEAR_MAX]:
        d = r["total"] - PASS_LINE
        rank_txt = ("%d / %d" % (r["class_rank"], n)) if r["class_rank"] else "缺考不计名次"
        watch.append(["已过线" if d >= 0 else "未过线", r["name"], fg(r["total"]),
                      signed(d), rank_txt])

    below = [nm for nm in SUBJ_NAMES
             if statistics.mean([r["scores"][nm] for r in ok]) < gidx["subjects"][nm]["mean"]]
    notes = ["临界生口径：总分距及格线 %d 分 ±%d 分以内，按离线的远近排，最多列 %d 人（本班共 %d 人落在这一带）。"
             % (PASS_LINE, NEAR_BAND, NEAR_MAX, len(near)),
             "及格率 = 总分 ≥ %d 分的人数 ÷ %d 人（有效人数，缺考剔除）；"
             "优秀率 = 总分 ≥ %d 分的人数 ÷ %d 人。" % (PASS_LINE, n, GOOD_LINE, n)]
    if len(below) == len(SUBJ_NAMES):
        notes.append("提醒：这里说的「最强」只是本班五科里最好的一科 —— 五科其实全部低于年级均分"
                     "（最接近的是%s，差 %s 分）。")
        notes[-1] = notes[-1] % (strong, f1(gidx["subjects"][strong]["mean"]
                                            - statistics.mean([r["scores"][strong] for r in ok])))
    elif below:
        notes.append("五科里有 %d 科低于年级均分：%s。"
                     % (len(below), "、".join(below)))
    if zero:
        all_totals = [r["total"] for r in members]
        notes.append("人数口径：本班 %d 人，其中 %d 人总分 0（缺考/未录入），"
                     "均分、及格率、优秀率、标准差与名次都按其余 %d 人算；"
                     "把缺考算进去会把均分拉到 %s 分、标准差抬到 %s 分（那是假的）。"
                     % (len(members), len(zero), n, f1(statistics.mean(all_totals)),
                        f1(statistics.pstdev(all_totals))))
    _guard(verdict, threshold, *notes)

    return {
        "kind": "class",
        "title": "一页纸诊断 ｜ %s" % cls,
        "sub": ("%d 人 · 九年级月考 · 满分 %d 分 —— 给班主任看的一页"
                % (len(members) if not zero else len(members), TOTAL_FULL)
                + ("" if not zero else "（其中 %d 人缺考/未录入，统计按 %d 人算）"
                   % (len(members) - n, n))),
        "big": [("均分", f1(mean)), ("及格率", f1(pass_rate) + "%"),
                ("优秀率", f1(good_rate) + "%"), ("标准差", f1(sigma))],
        "verdict": verdict,
        "threshold": threshold,
        "strong_subject": "%s %s%%（年级 %s%%）" % (strong, f1(rates[strong]),
                                                  f1(rate(gidx["subjects"][strong]["mean"], FULL[strong]))),
        "weak_subject": "%s %s%%（年级 %s%%）" % (weak, f1(rates[weak]),
                                                f1(rate(gidx["subjects"][weak]["mean"], FULL[weak]))),
        "tables": [
            {"title": "五科对比", "head": ["学科", "班级均分", "满分", "班级得分率",
                                          "年级均分", "差", "年级得分率"],
             "rows": subject_rows},
            {"title": "要单独盯的人（临界生，按离线远近）",
             "head": ["状态", "姓名", "总分", "距及格线", "班级名次"], "rows": watch},
        ],
        "notes": notes,
        "file": "班级诊断_%s.html" % cls,
    }


# ------------------------------------------------------------------ #89 教师话术卡
def build_talk_card(rows, gidx, stu):
    cls = stu["cls"]
    rates = {nm: rate(stu["scores"][nm], FULL[nm]) for nm in SUBJ_NAMES}
    strong = max(SUBJ_NAMES, key=lambda nm: (rates[nm], FULL[nm]))
    weak = min(SUBJ_NAMES, key=lambda nm: (rates[nm], FULL[nm]))
    weak_gap = gidx["subjects"][weak]["mean"] - stu["scores"][weak]
    weak_pp = rate(gidx["subjects"][weak]["mean"], FULL[weak]) - rates[weak]
    spread = rates[strong] - rates[weak]
    dmini = stu["mini_rank"] - stu["grade_rank"]
    passed = stu["total"] >= PASS_LINE

    # 本班有效人数不足群组下限时，班内名次没有意义（1 人班也有"第 1 名"）：一对一话术卡里不报
    rank_txt = ("班级第 %d 名（本班 %d 人计入排名，共 %d 人）。"
                % (stu["class_rank"], stu["class_valid"], stu["class_size"])
                if stu["class_valid"] >= MIN_GROUP
                else "本班有效成绩只有 %d 人（少于群组下限 %d 人）：不排名次，也不做横向比较。"
                     % (stu["class_valid"], MIN_GROUP))
    facts = [
        "总分 %s 分（满分 %d），年级第 %d 名（共 %d 人计入排名，全年级 %d 人、其中 %d 人缺考不计），%s"
        % (fg(stu["total"]), TOTAL_FULL, stu["grade_rank"], gidx["n"], gidx["n_all"], gidx["n_absent"],
           rank_txt),
        "五科得分率：%s。" % "、".join("%s %s%%" % (nm, f1(rates[nm])) for nm in SUBJ_NAMES),
        "和年级均分 %s 分比，%s 分；及格线 %d 分，%s。"
        % (f1(gidx["mean"]), ("高 %s" % f1(stu["total"] - gidx["mean"])) if stu["total"] >= gidx["mean"]
           else ("低 %s" % f1(gidx["mean"] - stu["total"])), PASS_LINE,
           "已经过线" if passed else ("还差 %s 分" % f1(PASS_LINE - stu["total"]))),
    ]

    if weak_gap > 0:
        reason_weak = ("最弱的是%s：%s 分（得分率 %s%%），年级均分 %s 分，差 %s 分、%s 个百分点。"
                       "这一科就是总分没能再往上的位置。"
                       % (weak, fg(stu["scores"][weak]), f1(rates[weak]),
                          f1(gidx["subjects"][weak]["mean"]), f1(weak_gap), f1(weak_pp)))
    else:
        reason_weak = ("最弱的是%s：得分率 %s%%，比年级均分（%s%%）还高 %s 个百分点 —— "
                       "不是掉队科，是目前的天花板在这一科。"
                       % (weak, f1(rates[weak]), f1(rate(gidx["subjects"][weak]["mean"], FULL[weak])),
                          f1(-weak_pp)))
    reasons = [
        reason_weak,
        "五科得分率最高 %s %s%%、最低 %s %s%%，极差 %s 个百分点（判定阈值 %s 个百分点），%s。"
        % (strong, f1(rates[strong]), weak, f1(rates[weak]), f1(spread), fg(BIAS_PP),
           "偏科很明显" if spread >= BIAS_PP else "各科比较均衡"),
        "语数外小总分年级第 %d 名，总分年级第 %d 名：物理化学把名次%s %d 名。"
        % (stu["mini_rank"], stu["grade_rank"], "往前推了" if dmini > 0 else ("往后拉了" if dmini < 0 else "没改变"),
           abs(dmini)),
    ]

    # 下一步：动作要能落到下周就能做的事
    if weak_gap > 0:
        target = gidx["subjects"][weak]["mean"]
        gain = target - stu["scores"][weak]
    else:
        target = stu["scores"][weak] + FULL[weak] * 0.05
        gain = FULL[weak] * 0.05
    would = stu["total"] + gain
    nexts = [
        "%s：下周起每周 2 次、每次 20 分钟，只做同一种题型（先 10 道），错一道当场重做一道。" % weak,
        "家长这边只做一件事：每次练习掐表 20 分钟、把错的那道抄在纸上，别讲方法 —— 讲方法由老师来。",
        "下次月考的验收指标：%s 得分率达到 %s%%（该科从 %s 分提到 %s 分），总分从 %s 提到 %s 以上，"
        "年级名次从第 %d 名进到第 %d 名。"
        % (weak, f1(rate(target, FULL[weak])), fg(stu["scores"][weak]), f1(target),
           fg(stu["total"]), f1(would), stu["grade_rank"], rank_if(rows, would)),
    ]

    # 不要说的话：三条，每条都点明「为什么没有数据支撑」
    dont = [
        ("「孩子就是不认真／态度有问题」",
         "成绩单里只有分数，没有一个字段记录过学习态度。这次%s只拿了 %s 分（满分 %s），"
         "能确认的是这一科的得分率 %s%%，不能确认的是原因。把它说成态度，家长回去只会骂一顿，"
         "那个题型还是没人补。" % (weak, fg(stu["scores"][weak]), fg(FULL[weak]), f1(rates[weak]))),
        ("「再这样下去就考不上好学校了」",
         "这一次月考总分 %s 分、年级第 %d 名，只说明这一次考试。没有任何一次月考能推出升学结果，"
         "说出来只会让家长恐慌、让孩子泄气。" % (fg(stu["total"]), stu["grade_rank"])),
        ("「人家能考 %s 分，你怎么才考 %s 分」" % (fg(gidx["max"]), fg(stu["total"])),
         "年级最高分 %s 分和这名学生的起点没有可比性。他真正的对照线是及格线 %d 分"
         "（现在%s）和年级均分 %s 分（相差 %s 分），拿第一名当尺子只会让人放弃。"
         % (fg(gidx["max"]), PASS_LINE, "已过" if passed else "还差 %s 分" % f1(PASS_LINE - stu["total"]),
            f1(gidx["mean"]), f1(abs(stu["total"] - gidx["mean"])))),
    ]
    _guard(*facts, *reasons, *nexts)

    return {
        "kind": "talk",
        "title": "教师话术卡 ｜ %s %s" % (cls, stu["name"]),
        "sub": "跟家长通电话/见面之前看一眼这一页 · 全部数字来自本次月考成绩单",
        "facts": facts,
        "reasons": reasons,
        "nexts": nexts,
        "dont": dont,
        "file": "话术卡_%s.html" % stu["name"],
    }


# ------------------------------------------------------------------ 纯文本版（打印到屏幕上）
def card_text(card):
    out = []
    line = "=" * 62
    out.append(line)
    out.append(card["title"])
    out.append(card["sub"])
    out.append(line)
    if card["kind"] in ("student", "class"):
        out.append("  " + " ｜ ".join("%s %s" % (k, v) for k, v in card["big"]))
        out.append("")
    if card["kind"] == "class":
        out.append("一句话结论：" + card["verdict"])
        out.append("")
        out.append(card["threshold"])
        out.append("")
        out.append("最强学科：%s" % card["strong_subject"])
        out.append("最弱学科：%s" % card["weak_subject"])
        out.append("")
    if card["kind"] == "student":
        out.append("最强科：%s" % card["strong"])
        out.append("最弱科：%s（极差 %s）" % (card["weak"], card["spread"]))
        out.append("")
    for t in (card.get("tables") or []):
        if t.get("title"):
            out.append("【%s】" % t["title"])
        widths = []
        for i, h in enumerate(t["head"]):
            w = max([_w(h)] + [_w(r[i]) for r in t["rows"]])
            widths.append(w)
        def fmt_row(vals):
            return "  " + "  ".join(_pad(v, widths[i], right=(i > 0))
                                    for i, v in enumerate(vals))
        out.append(fmt_row(t["head"]))
        out.append("  " + "-" * (sum(widths) + 2 * (len(widths) - 1)))
        for r in t["rows"]:
            out.append(fmt_row(r))
        out.append("")
    if card["kind"] == "student":
        out.append("【三条建议】")
        for i, a in enumerate(card["advices"], 1):
            out.append("%d. %s" % (i, a))
        out.append("")
    if card["kind"] == "talk":
        out.append("【一、先说事实】")
        for i, x in enumerate(card["facts"], 1):
            out.append("%d. %s" % (i, x))
        out.append("")
        out.append("【二、再说原因】")
        for i, x in enumerate(card["reasons"], 1):
            out.append("%d. %s" % (i, x))
        out.append("")
        out.append("【三、最后说下一步】")
        for i, x in enumerate(card["nexts"], 1):
            out.append("%d. %s" % (i, x))
        out.append("")
        out.append("【四、不要说的话】")
        for i, (say, why) in enumerate(card["dont"], 1):
            out.append("%d. %s" % (i, say))
            out.append("   为什么不能说：%s" % why)
        out.append("")
    if card.get("notes"):
        out.append("【口径】")
        for x in card["notes"]:
            out.append("· " + x)
    return "\n".join(out)


# ------------------------------------------------------------------ HTML（A4、零外链、打印隐藏按钮、手机单列）
CSS = """
*{box-sizing:border-box}
body{margin:0;padding:14px;background:#f3f4f6;color:#1a1a1a;
     font:14px/1.65 "Microsoft YaHei","PingFang SC","Hiragino Sans GB",sans-serif}
.page{background:#fff;max-width:780px;margin:0 auto;padding:20px 22px;border:1px solid #e5e5e5;border-radius:8px}
.toolbar{max-width:780px;margin:0 auto 10px;text-align:right}
button{font:inherit;padding:7px 16px;border:1px solid #1f4e79;border-radius:6px;
       background:#1f4e79;color:#fff;cursor:pointer}
h1{font-size:20px;margin:0 0 3px;color:#1f4e79}
.sub{color:#666;font-size:12px;margin-bottom:10px}
.big{display:grid;grid-template-columns:repeat(4,1fr);gap:8px;margin:12px 0}
.big div{border:1px solid #e3e7ec;border-radius:7px;padding:9px 6px;text-align:center;background:#fafbfc}
.big b{display:block;font-size:21px;color:#1f4e79;line-height:1.25}
.big span{font-size:12px;color:#666}
.verdict{border-left:5px solid #c0392b;background:#fdf3f2;padding:10px 12px;margin:12px 0;
         font-size:15px;font-weight:600;line-height:1.6}
.thr{font-size:12px;color:#555;background:#f6f8fa;border:1px dashed #cfd8e3;padding:8px 10px;
     border-radius:6px;margin-bottom:12px}
.two{display:grid;grid-template-columns:1fr 1fr;gap:10px;margin:10px 0}
.pill{border:1px solid #e3e7ec;border-radius:7px;padding:9px 11px;background:#fafbfc}
.pill em{display:block;font-style:normal;font-size:12px;color:#666}
.pill b{font-size:15px;color:#1f4e79}
table{width:100%;border-collapse:collapse;font-size:13px;margin:6px 0 2px}
caption{text-align:left;font-size:13px;font-weight:600;color:#333;padding:8px 0 4px}
th,td{border-bottom:1px solid #e8ebef;padding:5px 4px;text-align:right;white-space:nowrap}
th{background:#f2f6fa;font-weight:600;color:#33475b}
th:first-child,td:first-child,th:nth-child(2),td:nth-child(2){text-align:left}
ol{margin:6px 0 6px 20px;padding:0}
ol li{margin:7px 0}
.dont{border:1px solid #f0c8c4;background:#fff8f7;border-radius:8px;padding:4px 14px 10px;margin:10px 0}
.dont h2{font-size:15px;color:#b1372b;margin:10px 0 4px}
.dont .say{font-weight:700;color:#b1372b;margin:8px 0 2px}
.dont .why{color:#444;font-size:13px}
h2{font-size:15px;color:#1f4e79;margin:16px 0 4px;padding-bottom:3px;border-bottom:1px solid #e3e7ec}
.notes{font-size:12px;color:#555;margin-top:14px;border-top:1px dashed #dfe4ea;padding-top:8px}
.notes div{margin:3px 0}
.blank{margin-top:16px;border:1px solid #c9d3de;border-radius:4px;padding:10px 12px;background:#fbfcfe}
.blank .bt{font-size:13px;font-weight:700;color:#1f4e79;margin-bottom:8px}
.blank .bl{font-size:12px;color:#555;margin:6px 0 2px}
.blank .bline{border-bottom:1px solid #b9c3ce;height:20px}
@media print{
  body{background:#fff;padding:0;font-size:12.5px}
  .toolbar{display:none!important}
  .page{max-width:none;border:none;border-radius:0;padding:0;margin:0}
  .no-print{display:none!important}
  table{page-break-inside:avoid}
  h2,.big,table{break-inside:avoid}
}
@media screen and (max-width:480px){
  body{padding:8px;font-size:13px}
  .page{padding:14px 12px}
  .big{grid-template-columns:1fr}
  .two{grid-template-columns:1fr}
  table{font-size:12px}
  th,td{padding:4px 2px}
  .toolbar{text-align:center}
}
@page{size:A4;margin:12mm}
"""


def _html_table(t):
    head = "".join("<th>%s</th>" % html.escape(str(h)) for h in t["head"])
    body = "".join("<tr>%s</tr>" % "".join("<td>%s</td>" % html.escape(str(c)) for c in row)
                   for row in t["rows"])
    cap = "<caption>%s</caption>" % html.escape(t["title"]) if t.get("title") else ""
    return "<table>%s<thead><tr>%s</tr></thead><tbody>%s</tbody></table>" % (cap, head, body)


def card_html(card):
    e = html.escape
    parts = []
    parts.append("<!doctype html>")
    parts.append('<html lang="zh-CN"><head><meta charset="utf-8">')
    parts.append('<meta name="viewport" content="width=device-width, initial-scale=1">')
    parts.append("<title>%s</title>" % e(card["title"]))
    parts.append("<style>%s</style>" % CSS)
    parts.append("</head><body>")
    parts.append('<div class="toolbar"><button type="button" onclick="window.print()">打印这一页（A4）</button></div>')
    parts.append('<div class="page">')
    parts.append("<h1>%s</h1>" % e(card["title"]))
    parts.append('<div class="sub">%s</div>' % e(card["sub"]))

    if card["kind"] in ("student", "class"):
        parts.append('<div class="big">%s</div>' % "".join(
            "<div><b>%s</b><span>%s</span></div>" % (e(v), e(k)) for k, v in card["big"]))
    if card["kind"] == "class":
        parts.append('<div class="verdict">一句话结论：%s</div>' % e(card["verdict"]))
        parts.append('<div class="thr">%s</div>' % e(card["threshold"]))
        parts.append('<div class="two">')
        parts.append('<div class="pill"><em>最强学科（按得分率）</em><b>%s</b></div>' % e(card["strong_subject"]))
        parts.append('<div class="pill"><em>最弱学科（按得分率）</em><b>%s</b></div>' % e(card["weak_subject"]))
        parts.append("</div>")
    if card["kind"] == "student":
        parts.append('<div class="two">')
        parts.append('<div class="pill"><em>最强科（按得分率）</em><b>%s</b></div>' % e(card["strong"]))
        parts.append('<div class="pill"><em>最弱科（按得分率）</em><b>%s</b>　极差 %s</div>'
                     % (e(card["weak"]), e(card["spread"])))
        parts.append("</div>")
    for t in (card.get("tables") or []):
        parts.append(_html_table(t))
    if card["kind"] == "student":
        parts.append("<h2>三条建议（每条都带数字，能当场对账）</h2>")
        parts.append("<ol>%s</ol>" % "".join("<li>%s</li>" % e(a) for a in card["advices"]))
    if card["kind"] == "talk":
        parts.append("<h2>一、先说事实（全部是本页数字）</h2>")
        parts.append("<ol>%s</ol>" % "".join("<li>%s</li>" % e(a) for a in card["facts"]))
        parts.append("<h2>二、再说原因（依据在同一行里）</h2>")
        parts.append("<ol>%s</ol>" % "".join("<li>%s</li>" % e(a) for a in card["reasons"]))
        parts.append("<h2>三、最后说下一步（具体到动作）</h2>")
        parts.append("<ol>%s</ol>" % "".join("<li>%s</li>" % e(a) for a in card["nexts"]))
        parts.append('<div class="dont"><h2>四、不要说的话（这三句请咽回去）</h2>')
        for say, why in card["dont"]:
            parts.append('<div class="say">✗ %s</div><div class="why">为什么不能说：%s</div>' % (e(say), e(why)))
        parts.append("</div>")
    # 这一块是故意留空的：数据只描述了学生的一部分，剩下的必须由老师来填。
    # 不留这一块，一份写满"该补哪科"的纸会把一个孩子说尽。
    parts.append('<div class="blank"><div class="bt">数据看不到的部分（请老师填）</div>'
                 '<div class="bl">这个学生在班里承担着什么：</div><div class="bline"></div>'
                 '<div class="bl">他擅长、但分数体现不出来的：</div><div class="bline"></div></div>')
    if card.get("notes"):
        parts.append('<div class="notes"><div><b>口径</b></div>%s</div>'
                     % "".join("<div>· %s</div>" % e(x) for x in card["notes"]))
    parts.append("</div></body></html>")
    return "\n".join(parts)


def write_card(card):
    os.makedirs(OUT_DIR, exist_ok=True)
    path = os.path.join(OUT_DIR, card["file"])
    with open(path, "w", encoding="utf-8") as f:
        f.write(card_html(card))
    return path


# ------------------------------------------------------------------ CLI
def main(argv=None):
    ap = argparse.ArgumentParser(
        description="教师工具 · 一人一页学情卡 / 一页纸班级诊断 / 教师话术卡（0 次模型调用）")
    ap.add_argument("--xlsx", default=XLSX, help="成绩单路径（默认 %s）" % XLSX)
    sub = ap.add_subparsers(dest="cmd")

    p1 = sub.add_parser("student", help="一人一页学情卡")
    p1.add_argument("--name", required=True, help="学生姓名")
    p1.add_argument("--class", dest="cls", default=None, help="重名时用来指定班级")

    p2 = sub.add_parser("class", help="一页纸班级诊断")
    p2.add_argument("--name", required=True, help="班级名，例如 九年级2班")
    p2.add_argument("--min-group", dest="min_group", type=int, default=MIN_GROUP,
                    help="群组下限：有效人数低于它拒绝生成班级诊断（默认 %d）" % MIN_GROUP)

    p3 = sub.add_parser("talk", help="教师话术卡")
    p3.add_argument("--name", required=True, help="学生姓名")
    p3.add_argument("--class", dest="cls", default=None, help="重名时用来指定班级")

    args = ap.parse_args(argv)
    if not args.cmd:
        ap.print_help()
        return 0

    try:
        rows = with_ranks(load_rows(args.xlsx))
        gidx = grade_index(rows)
        if args.cmd == "student":
            card = build_student_card(rows, gidx, pick_student(rows, args.name, getattr(args, "cls", None)))
        elif args.cmd == "class":
            card = build_class_card(rows, gidx, args.name, getattr(args, "min_group", MIN_GROUP))
        else:
            card = build_talk_card(rows, gidx, pick_student(rows, args.name, getattr(args, "cls", None)))
    except UserError as e:
        print("！" + str(e))
        return 2

    path = write_card(card)
    print(card_text(card))
    print("")
    print(line_sep())
    print("HTML 已生成：%s" % path)
    print("打印：浏览器打开这一页 → Ctrl+P → A4（按钮在打印时自动隐藏）")
    return 0


def line_sep():
    return "-" * 62


if __name__ == "__main__":
    sys.exit(main())
