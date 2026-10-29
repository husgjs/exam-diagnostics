# -*- coding: utf-8 -*-
"""ingest.py 的验收：用四类真实会遇到的烂表来测。"""
import io, os, sys, tempfile
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl import Workbook
import ingest

PASS, FAIL = 0, 0
HERE = os.path.dirname(os.path.abspath(__file__))

def check(name, cond, extra=""):
    global PASS, FAIL
    if cond:
        PASS += 1; print("  PASS  " + name)
    else:
        FAIL += 1; print("  FAIL  " + name + ("   " + str(extra) if extra else ""))

def tmpfile():
    os.makedirs(os.path.join(HERE, "_tmp"), exist_ok=True)
    fd, p = tempfile.mkstemp(suffix=".xlsx", dir=os.path.join(HERE, "_tmp"))
    os.close(fd); return p

print("== 1. 多 sheet：两个班各一个表，都要读到 ==")
p = tmpfile()
wb = Workbook(); ws1 = wb.active; ws1.title = "一班"
ws1.append(["姓名", "班级", "总分"]); ws1.append(["张三", "一班", 480]); ws1.append(["李四", "一班", 455])
ws2 = wb.create_sheet("二班"); ws2.append(["姓名", "班级", "总分"]); ws2.append(["王五", "二班", 500])
wb.save(p)
recs, rep = ingest.read_all_sheets(p)
check("两个 sheet 都被扫到", len(rep["sheets"]) == 2, rep["sheets"])
check("两个 sheet 都读出了数据", len(rep["used"]) == 2, rep["used"])
check("一共 3 条记录", len(recs) == 3, len(recs))
check("记录里带 sheet 名", {r["__sheet"] for r in recs} == {"一班", "二班"})
check("姓名与总分都对", recs[0]["name"] == "张三" and recs[0]["total"] == 480, recs[0])
os.remove(p)

print("\n== 2. 表头不在第一行 ==")
p = tmpfile()
wb = Workbook(); ws = wb.active; ws.title = "成绩"
ws.append(["九年级第一次月考成绩统计表"]); ws.append([]); ws.append(["制表：教务处", None, None])
ws.append(["学号", "姓名", "总分"]); ws.append(["1001", "赵六", 512])
wb.save(p)
recs, rep = ingest.read_all_sheets(p)
check("认出了表头（跳过了前 3 行）", len(recs) == 1, len(recs))
check("取到姓名", recs and recs[0]["name"] == "赵六", recs)
check("取到学号", recs and recs[0]["sid"] == "1001", recs)
os.remove(p)

print("\n== 3. 合并单元格 ==")
p = tmpfile()
wb = Workbook(); ws = wb.active; ws.title = "逐题"
ws.append(["学号", "姓名", "1", "2", "3"])
ws.append(["1001", "钱七", 3, 0, 3]); ws.append([None, None, 2, 2, 3])
ws.merge_cells("A2:A3"); ws.merge_cells("B2:B3")
wb.save(p)
rows = ingest.sheet_rows(p, "逐题")
check("合并单元格被填满（第二行也有学号）", rows[2][0] == "1001", rows[2][:2])
check("第二行也有姓名", rows[2][1] == "钱七", rows[2][:2])
os.remove(p)

print("\n== 4. 宽表 -> 长表 ==")
hdr = ["学号", "姓名", "1", "2", "第3题", "Q4", "总分"]
qcols = ingest.detect_question_cols(hdr)
check("认出 4 个题号列", len(qcols) == 4, qcols)
check("「总分」不被当成题号", all(hdr[i] != "总分" for i, _ in qcols), qcols)
rows = [["1001", "孙八", 3, 0, 2, 3, 8]]
long_rows, missing = ingest.wide_to_long(hdr, rows, student_key=1, full_map={"1": 3, "2": 3, "3": 3, "4": 3})
check("转出 4 行长表", len(long_rows) == 4, len(long_rows))
check("得分为 0 也保留（不等于没做）", any(r["got"] == 0 for r in long_rows))
check("学生标识正确", all(r["student"] == "孙八" for r in long_rows))
check("满分齐了，missing 为空", missing == [], missing)

print("\n== 5. 满分缺失必须报出来 ==")
long_rows, missing = ingest.wide_to_long(hdr, rows, student_key=1, full_map={"1": 3})
check("缺满分的题号被列出", set(missing) == {"2", "3", "4"}, missing)
check("缺满分的行 full 为 None（没编数值）", all(r["full"] is None for r in long_rows if r["qno"] in missing))

print("\n== 6. 认不出表头时不许硬算 ==")
p = tmpfile()
wb = Workbook(); ws = wb.active
for row in [["这是说明"], ["没有表头"], ["随便写点"]]: ws.append(row)
wb.save(p)
recs, rep = ingest.read_all_sheets(p)
check("认不出表头 -> 0 条记录", len(recs) == 0, len(recs))
check("记进 skipped 并写明原因", len(rep["skipped"]) == 1 and "找不到" in rep["skipped"][0][1], rep["skipped"])
os.remove(p)

print("\n== 7. 同义列名 ==")
for variant in ["姓名", "学生姓名", "学员姓名", "学生"]:
    check("认得出「%s」" % variant, ingest.canon(["学号", variant]).get("name") == 1)
check("全角空格也能认", ingest.canon(["姓　名"]).get("name") == 0)

print()
print("RESULT:", "ALL PASS" if FAIL == 0 else "HAS FAILURE", "| 通过 %d，失败 %d" % (PASS, FAIL))
sys.exit(0 if FAIL == 0 else 1)
