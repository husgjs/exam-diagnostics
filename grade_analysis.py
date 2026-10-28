
# -*- coding: utf-8 -*-
"""年级成绩单分析 —— 适配「班级 + 姓名 + 各科分数 + 总分 + 名次」这种真实表。

用法：
    python grade_analysis.py 全年级成绩单.xlsx
    python grade_analysis.py 全年级成绩单.xlsx --full 语文=100 数学=150 英语=150 物理=100 化学=100
"""
import io, os, sys, re, argparse
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from collections import defaultdict, OrderedDict

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
WARN = PatternFill("solid", fgColor="FDE2E2")
GOOD = PatternFill("solid", fgColor="E4F2E4")
TITLE = Font(bold=True, size=13, color="1F4E79")

# 常见满分，表里没给时按这个猜；可用 --full 覆盖
DEFAULT_FULL = {"语文": 100, "数学": 150, "英语": 150, "物理": 100, "化学": 100,
                "政治": 100, "历史": 100, "生物": 100, "地理": 100, "体育": 50}

def to_num(v):
    if v is None: return None
    if isinstance(v, (int, float)): return float(v)
    s = str(v).strip()
    if s in ("", "缺考", "缺", "-", "—", "/"): return None
    try: return float(re.sub(r"[^0-9.\-]", "", s))
    except Exception: return None

def head_style(ws, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEAD_FILL; cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"

def cls_order(name):
    m = re.search(r"(\d+)", str(name))
    return int(m.group(1)) if m else 99

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--full", nargs="*", default=[],
                    help="各科满分，如 数学=150 语文=100")
    a = ap.parse_args()

    full_map = dict(DEFAULT_FULL)
    for kv in a.full:
        if "=" in kv:
            k, v = kv.split("=", 1)
            full_map[k.strip()] = float(v)

    ws = load_workbook(a.file, data_only=True).active
    head = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    print("表头:", head)

    def find(*keys):
        for i, h in enumerate(head):
            if any(k in h for k in keys): return i
        return None

    ci_cls = find("班级")
    ci_name = find("姓名")
    # 总分要精确匹配：表里可能同时有「小总分」（语数英小计）和「总分」，
    # 用包含匹配会先命中「小总分」，整份统计就全错了。
    ci_total = next((i for i, h in enumerate(head) if h.strip() == "总分"), None)
    if ci_total is None:
        ci_total = next((i for i, h in enumerate(head) if "总分" in h and "小" not in h), None)
    ci_rank = next((i for i, h in enumerate(head) if h.strip() == "总名次"), None)
    if ci_rank is None:
        ci_rank = find("名次")
    # 科目列：以「分数」结尾且不是总分/小总分
    ci_subj = [(i, h.replace("分数", "")) for i, h in enumerate(head)
               if h.endswith("分数") and "总" not in h]
    print("班级列:", ci_cls, "｜姓名列:", ci_name, "｜总分列:", ci_total, "｜名次列:", ci_rank)
    print("科目:", [s[1] for s in ci_subj])

    rows = []
    for r in range(2, ws.max_row + 1):
        vals = [c.value for c in ws[r]]
        if all(v is None or str(v).strip() == "" for v in vals): continue
        nm = vals[ci_name] if ci_name is not None else ""
        if not nm: continue
        rec = {"cls": str(vals[ci_cls]).strip() if ci_cls is not None else "全体",
               "name": str(nm).strip(),
               "subj": {}, "total": to_num(vals[ci_total]) if ci_total is not None else None,
               "rank": to_num(vals[ci_rank]) if ci_rank is not None else None}
        for i, s in ci_subj:
            v = to_num(vals[i]) if i < len(vals) else None
            if v is not None: rec["subj"][s] = v
        if rec["total"] is None and rec["subj"]:
            rec["total"] = sum(rec["subj"].values())
        rows.append(rec)
    # 总分 0 = 缺考/未录入，算统计前必须先剔除；否则几个缺考就能造出假的"两极分化"
    _zeros = [d for d in rows if (d.get("total") or 0) <= 0]
    if _zeros:
        print("数据质检：剔除 %d 人总分 0（缺考/未录入）：%s"
              % (len(_zeros), "、" .join("%s(%s)" % (d["name"], d["cls"]) for d in _zeros[:10])))
        rows = [d for d in rows if (d.get("total") or 0) > 0]
    print("学生 %d 人" % len(rows))
    if not rows:
        raise SystemExit("表里没有有效学生数据：请检查「姓名」列是否填了、总分列是否存在")
    if len(rows) < 2:
        raise SystemExit("只有 1 名学生的数据，做不了班级/年级分析（至少需要 2 人）")

    classes = sorted({r["cls"] for r in rows}, key=cls_order)
    subjects = [s for _, s in ci_subj]
    full_total = sum(full_map.get(s, 100) for s in subjects)
    print("班级 %d 个 ｜ 科目 %d 科 ｜ 估算满分 %.0f" % (len(classes), len(subjects), full_total))
    print()

    totals = [r["total"] for r in rows if r["total"] is not None]
    if not totals:
        raise SystemExit("没有任何有效总分，无法分析")
    avg = sum(totals) / len(totals)
    PL, GL = full_total * 0.6, full_total * 0.85

    out = os.path.splitext(a.file)[0] + "_分析.xlsx"
    wb = Workbook()

    # ---- 1 班级对比 ----
    ws1 = wb.active; ws1.title = "班级对比"
    ws1.append(["班级", "人数", "总分均分", "最高", "最低", "及格率", "优秀率"] +
               ["%s均分" % s for s in subjects])
    head_style(ws1, 7 + len(subjects))
    by_cls = defaultdict(list)
    for r in rows: by_cls[r["cls"]].append(r)
    print("%-10s %4s %8s %7s %7s %7s %7s  %s" % ("班级", "人数", "总分均分", "最高", "最低", "及格率", "优秀率", "各科均分"))
    for c in classes:
        g = by_cls[c]
        ts = [x["total"] for x in g if x["total"] is not None]
        if not ts: continue
        pa = 100 * sum(1 for t in ts if t >= PL) / len(ts)
        gd = 100 * sum(1 for t in ts if t >= GL) / len(ts)
        subj_avg = []
        for s in subjects:
            vs = [x["subj"][s] for x in g if s in x["subj"]]
            subj_avg.append(round(sum(vs) / len(vs), 1) if vs else "")
        ws1.append([c, len(g), round(sum(ts) / len(ts), 1), max(ts), min(ts),
                    "%.0f%%" % pa, "%.0f%%" % gd] + subj_avg)
        print("%-10s %4d %8.1f %7.0f %7.0f %6.0f%% %6.0f%%  %s" % (
            c, len(g), sum(ts) / len(ts), max(ts), min(ts), pa, gd, subj_avg))

    # ---- 2 各科分析 ----
    ws2 = wb.create_sheet("各科分析")
    ws2.append(["科目", "满分", "年级均分", "得分率", "最高", "最低", "及格率", "判定"])
    head_style(ws2, 8)
    print()
    print("%-6s %5s %8s %7s %7s %7s %7s  %s" % ("科目", "满分", "均分", "得分率", "最高", "最低", "及格率", "判定"))
    for s in subjects:
        vs = [r["subj"][s] for r in rows if s in r["subj"]]
        if not vs: continue
        F = full_map.get(s, max(vs))
        rate = sum(vs) / len(vs) / F * 100
        pr = 100 * sum(1 for v in vs if v >= F * 0.6) / len(vs)
        judge = "偏难" if rate < 60 else ("偏易" if rate > 85 else "正常")
        ws2.append([s, F, round(sum(vs) / len(vs), 1), "%.0f%%" % rate, max(vs), min(vs),
                    "%.0f%%" % pr, judge])
        print("%-6s %5.0f %8.1f %6.0f%% %7.0f %7.0f %6.0f%%  %s" % (
            s, F, sum(vs) / len(vs), rate, max(vs), min(vs), pr, judge))

    # ---- 3 全年级分数段 ----
    step = 50
    segs = []
    for k in range(int(full_total // step)):
        lo = step * k; hi = step * (k + 1)
        segs.append((lo, hi, sum(1 for t in totals if lo < t <= hi)))
    segs.reverse()
    ws3 = wb.create_sheet("分数段分布")
    ws3.append(["总分段", "人数", "占比", "图示"]); head_style(ws3, 4)
    print()
    print("== 全年级总分分布（均分 %.1f / %.0f）==" % (avg, full_total))
    for lo, hi, cnt in segs:
        if cnt == 0: continue
        ws3.append(["%d–%d" % (lo, hi), cnt, "%.0f%%" % (100 * cnt / len(totals)), "█" * min(cnt, 60)])
        print("   %3d–%3d  %3d 人  %s" % (lo, hi, cnt, "█" * min(cnt, 60)))

    # ---- 4 学生明细 ----
    ws4 = wb.create_sheet("学生明细")
    ws4.append(["年级名次", "班级", "姓名"] + subjects + ["总分", "班级名次"])
    head_style(ws4, 5 + len(subjects))
    rows_sorted = sorted(rows, key=lambda x: -(x["total"] or 0))
    cls_rank = defaultdict(int)
    for r in rows_sorted:
        cls_rank[r["cls"]] += 1
        r["cls_rank"] = cls_rank[r["cls"]]
        ws4.append([r["rank"] or "", r["cls"], r["name"]] +
                   [r["subj"].get(s, "") for s in subjects] + [r["total"], r["cls_rank"]])
        mark = WARN if (r["total"] or 0) < PL else (GOOD if (r["total"] or 0) >= GL else None)
        if mark: ws4.cell(row=ws4.max_row, column=4 + len(subjects)).fill = mark
    # 缺考的不进统计，但打上「缺」字留在表里
    for d in _zeros:
        ws4.append(["缺考", d["cls"], d["name"]] + ["缺"] * len(subjects) + ["缺考", "缺考"])
        for _c in range(1, 6 + len(subjects)):
            ws4.cell(row=ws4.max_row, column=_c).fill = WARN
    for i, w in enumerate([9, 11, 10] + [8] * len(subjects) + [8, 9], 1):
        ws4.column_dimensions[get_column_letter(i)].width = w
    ws4.auto_filter.ref = "A1:%s%d" % (get_column_letter(5 + len(subjects)), ws4.max_row)

    wb.save(out)
    print()
    print("已生成:", out)

if __name__ == "__main__":
    main()
