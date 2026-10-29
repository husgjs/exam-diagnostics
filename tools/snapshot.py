# -*- coding: utf-8 -*-
"""产物内容快照：用于证明「改读表方式没把结果改坏」。

为什么不能直接比文件哈希：xlsx 是 zip，里面含时间戳与压缩差异，
同样的数据两次生成字节也不同。所以比的是**单元格值序列**的摘要。

用法：
    python tools/snapshot.py --dir 某目录 --out _tmp/base.json
    python tools/snapshot.py --dir 某目录 --diff _tmp/base.json
"""
import argparse
import glob
import hashlib
import io
import json
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

try:
    import openpyxl
except ImportError:
    print("需要 openpyxl")
    sys.exit(3)


def cell_str(v):
    """值的规范化：避免 3 / 3.0 / '3' 造成假差异，但也不把不同值抹平。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        f = float(v)
        if f == int(f):
            return str(int(f))
        return repr(round(f, 6))
    return str(v).strip()


def file_digest(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    try:
        out = {}
        for name in wb.sheetnames:
            ws = wb[name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append("\x1f".join(cell_str(c) for c in row))
            blob = "\n".join(rows).rstrip("\x1f\n")
            out[name] = {
                "sha": hashlib.sha256(blob.encode("utf-8")).hexdigest()[:16],
                "rows": len(rows),
                "chars": len(blob),
            }
        return out
    finally:
        wb.close()


def scan(folder):
    res = {}
    for p in sorted(glob.glob(os.path.join(folder, "*.xlsx"))):
        if os.path.basename(p).startswith("~$"):
            continue
        try:
            res[os.path.basename(p)] = file_digest(p)
        except Exception as e:
            res[os.path.basename(p)] = {"__error__": str(e)[:120]}
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", required=True)
    ap.add_argument("--out", default=None)
    ap.add_argument("--diff", default=None, help="与已有快照比对")
    a = ap.parse_args()

    cur = scan(a.dir)
    if not cur:
        # 扫到 0 个 xlsx 时，「一致」是空口径，不能当证据 —— 直接失败
        print("空口径：%s 里没有任何 .xlsx，本次比对不成立" % a.dir)
        sys.exit(2)
    if a.diff:
        base = json.load(open(a.diff, encoding="utf-8"))
        bad = []
        for fn in sorted(set(base) | set(cur)):
            b, c = base.get(fn), cur.get(fn)
            if b == c:
                continue
            if b is None:
                bad.append((fn, "新增", "", ""))
                continue
            if c is None:
                bad.append((fn, "缺失", "", ""))
                continue
            for sh in sorted(set(b) | set(c)):
                if b.get(sh) != c.get(sh):
                    bad.append((fn, sh, str(b.get(sh, {}).get("sha", "-")), str(c.get(sh, {}).get("sha", "-"))))
        if bad:
            print("差异 %d 处：" % len(bad))
            for fn, sh, bs, cs in bad:
                print("  %-40s %-16s %s -> %s" % (fn[:40], sh[:16], bs[:12], cs[:12]))
            sys.exit(1)
        print("一致：%d 个文件，逐 sheet 内容摘要全等 ✓" % len(cur))
        return

    print("扫描 %s：%d 个 xlsx" % (a.dir, len(cur)))
    for fn, sheets in cur.items():
        print("  %-42s %s" % (fn[:42], ", ".join("%s=%s(%d行)" % (k, v.get("sha", "?")[:8], v.get("rows", 0)) for k, v in sheets.items())[:70]))
    if a.out:
        json.dump(cur, open(a.out, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        print("已写入", a.out)


if __name__ == "__main__":
    main()
